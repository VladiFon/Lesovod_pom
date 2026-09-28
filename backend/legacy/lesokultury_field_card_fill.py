# -*- coding: utf-8 -*-
"""Заполнение шаблонов полевых карточек лесных культур (по образцу
приложений 15/19 к постановлению Минлесхоза от 19.12.2016 N 80).

Общий хелпер для "Карточки перевода" (kartochka_perevoda_generator.py) и,
при необходимости в будущем, "Карточки инвентаризации" — обе карточки
собираются из одних и тех же структур данных (пробы + результаты по
породам + сверка с нормативом), поэтому таблицы-заготовки шаблона
заполняются одними и теми же функциями.

Растущие таблицы (число проб/пород заранее не известно) заполняются по
образцу uhody.py::build_osvetlenie_document — последняя строка-заготовка
клонируется столько раз, сколько не хватает под фактические данные."""
import copy

from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from akt_template_fill import set_cell_text  # переиспользуем авто-сжатие кегля

FONT = "Times New Roman"
BASE_SIZE = 11
MIN_SIZE = 8


def _clone_rows_for(table, n_extra_rows):
    """Клонирует последнюю строку таблицы n_extra_rows раз (глубокая копия
    XML строки — сохраняет форматирование/границы ячеек, как в
    uhody.py::build_osvetlenie_document). Возвращает None — таблица
    меняется на месте."""
    if n_extra_rows <= 0:
        return
    template_row_tr = table.rows[-1]._tr
    for _ in range(n_extra_rows):
        new_row_tr = copy.deepcopy(template_row_tr)
        table._tbl.append(new_row_tr)


def set_cell_text_wrap(table, row, col, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=BASE_SIZE):
    """Заполняет ячейку свободным (возможно многострочным) текстом БЕЗ
    авто-сжатия кегля — в отличие от set_cell_text (короткие структурные
    поля), здесь текст в ячейке нормально переносится по словам, как в
    обычном документе Word (для полей заключения комиссии — 12а-д)."""
    cell = table.rows[row].cells[col]
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    para = cell.paragraphs[0]
    for r in list(para.runs):
        r._element.getparent().remove(r._element)
    para.alignment = align

    text = "" if text is None else str(text)
    if text:
        run = para.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.bold = bold
    return cell


def fill_proba_table(table, proby):
    """Таблица пробных площадей (шаблон: заголовок + 1 строка-заготовка).
    proby — список объектов/словарей с полями nomer, razmer."""
    n = len(proby)
    if n == 0:
        set_cell_text(table, 1, 0, "")
        set_cell_text(table, 1, 1, "")
        return
    _clone_rows_for(table, n - 1)
    for i, p in enumerate(proby):
        nomer = p.get("nomer") if isinstance(p, dict) else p.nomer
        razmer = p.get("razmer") if isinstance(p, dict) else p.razmer
        set_cell_text(table, 1 + i, 0, nomer, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table, 1 + i, 1, razmer, align=WD_ALIGN_PARAGRAPH.CENTER)


def fill_rezultaty_table(table, rezultaty, itogo):
    """Таблица результатов обследования по породам (шаблон: заголовок +
    1 строка-заготовка + строка "Итого"). rezultaty — список объектов/
    словарей (poroda, vysazheno, prizhilos, srednyaya_vysota); itogo —
    словарь с ключами vysazheno/prizhilos/prizhivaemost_pct (как из
    _calc_prizhivaemost в app/routers/lesokultury.py)."""
    n = len(rezultaty)
    n_extra = max(n - 1, 0)
    _clone_rows_for(table, n_extra)
    # после клонирования строка "Итого" — всегда последняя
    total_row = len(table.rows) - 1
    for i, r in enumerate(rezultaty):
        get = (lambda k: r.get(k)) if isinstance(r, dict) else (lambda k: getattr(r, k, None))
        poroda = get("poroda")
        vysazheno = get("vysazheno")
        prizhilos = get("prizhilos")
        vysota = get("srednyaya_vysota")
        prizhivaemost = round(prizhilos / vysazheno * 100, 1) if vysazheno else 0
        row = 1 + i
        set_cell_text(table, row, 0, poroda, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(table, row, 1, vysazheno, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table, row, 2, prizhilos, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table, row, 3, f"{prizhivaemost}", align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table, row, 4, vysota, align=WD_ALIGN_PARAGRAPH.CENTER)

    set_cell_text(table, total_row, 0, "Итого", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(table, total_row, 1, itogo.get("vysazheno"), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table, total_row, 2, itogo.get("prizhilos"), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table, total_row, 3, itogo.get("prizhivaemost_pct"), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table, total_row, 4, "", align=WD_ALIGN_PARAGRAPH.CENTER)


def fill_normativ_table(table, normativ_check, glavnaya_poroda, tip_lesa,
                         kolichestvo_tys_na_ga, srednyaya_vysota_m):
    """Таблица сверки с нормативом (приложение 18) — 8 строк
    label/значение (шаблон kartochka_perevoda_shablon.docx, таблица 3).
    normativ_check — результат lesokultury_normativy.check_normativ_perevoda
    (может быть None, если порода не найдена в таблице норматива)."""
    nc = normativ_check or {}
    set_cell_text(table, 0, 1, glavnaya_poroda or "", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(table, 1, 1, tip_lesa or "", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(table, 2, 1, nc.get("trebuemoe_kolichestvo_tys_ga"), align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table, 3, 1, nc.get("trebuemaya_vysota_m"), align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table, 4, 1, kolichestvo_tys_na_ga, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table, 5, 1, srednyaya_vysota_m, align=WD_ALIGN_PARAGRAPH.CENTER)
    if nc.get("naideno"):
        soot = "Соответствует" if nc.get("sootvetstvuet") else "Не соответствует"
    else:
        soot = "Порода не найдена в таблице норматива (приложение 18)"
    set_cell_text(table, 6, 1, soot, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text_wrap(table, 7, 1, nc.get("predupredzhenie") or "")


def fill_zaklyuchenie_table(table, a="", b="", v="", g="", d=""):
    """Заключение комиссии (12 а-д) — шаблон, таблица 4 (5 строк, готовые
    подписи в столбце 0, свободный текст в столбце 1)."""
    for row, text in enumerate((a, b, v, g, d)):
        set_cell_text_wrap(table, row, 1, text)


def fill_signatures_table(table, predsedatel, chleny):
    """Подписи комиссии — шаблон, таблица 5 (3 строки: председатель + до
    2 членов). predsedatel/chleny — словари {dolzhnost, fio}."""
    if predsedatel:
        set_cell_text(table, 0, 1, predsedatel.get("dolzhnost", ""), align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(table, 0, 2, predsedatel.get("fio", ""), align=WD_ALIGN_PARAGRAPH.LEFT)
    for i, chlen in enumerate((chleny or [])[:2]):
        row = 1 + i
        set_cell_text(table, row, 1, chlen.get("dolzhnost", ""), align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(table, row, 2, chlen.get("fio", ""), align=WD_ALIGN_PARAGRAPH.LEFT)
