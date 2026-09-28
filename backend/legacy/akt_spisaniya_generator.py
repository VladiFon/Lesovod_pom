# -*- coding: utf-8 -*-
"""Генератор "Акта на списание погибших лесных культур" (приложение 20 к
постановлению Минлесхоза от 19.12.2016 N 80) — заполняет собственный
шаблон templates/akt_spisaniya_lesokultur_shablon.docx (вертикальная
таблица сведений об участке — акт печатается на один конкретный
списываемый участок, а не построчно на несколько, как в оригинале)."""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from akt_template_fill import set_cell_text
from lesokultury_field_card_fill import fill_signatures_table

_TEMPLATE_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "templates", "akt_spisaniya_lesokultur_shablon.docx",
)


def generate_akt_spisaniya(uchastok, latest_inventarizatsiya, output_path, **fields):
    """uchastok — dict из db.get_lesokultury_uchastok(); latest_inventarizatsiya —
    dict последней записи журнала с приживаемостью (db.list_lesokultury_meropriyatiya,
    отфильтрованная и отсортированная вызывающим кодом) либо None.

    fields (обязательные по акту — источника в базе нет):
        prichiny_gibeli, izrashodovano_tys_rub.
    fields (необязательные): reshenie_komissii, data_akta,
        predsedatel: {dolzhnost, fio}, chleny: [{dolzhnost, fio}, ...]."""
    template_path = fields.pop("template_path", None) or _TEMPLATE_PATH
    if not os.path.isfile(template_path):
        raise FileNotFoundError(f"Не найден шаблон акта на списание: {template_path}")

    doc = Document(template_path)
    tables = doc.tables
    if len(tables) < 2:
        raise ValueError("Шаблон акта на списание повреждён — ожидались 2 таблицы")
    t_info, t_podpisi = tables[:2]

    mestonahozhdenie = f"кв. {uchastok.get('kvartal') or '—'} / выд. {uchastok.get('vydel') or '—'}"
    prizhivaemost = latest_inventarizatsiya.get("prizhivaemost_pct") if latest_inventarizatsiya else None

    set_cell_text(t_info, 0, 1, mestonahozhdenie, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t_info, 1, 1, uchastok.get("god_sozdaniya") or "", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t_info, 2, 1, uchastok.get("ploshad") or "")
    set_cell_text(t_info, 3, 1, uchastok.get("glavnaya_poroda") or "", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t_info, 4, 1, uchastok.get("metod_sozdaniya") or "", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t_info, 5, 1, prizhivaemost if prizhivaemost is not None else "")
    set_cell_text(t_info, 6, 1, fields.get("prichiny_gibeli", ""), align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t_info, 7, 1, fields.get("reshenie_komissii", ""), align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t_info, 8, 1, fields.get("izrashodovano_tys_rub", ""))

    fill_signatures_table(t_podpisi, fields.get("predsedatel"), fields.get("chleny"))

    data_akta = fields.get("data_akta") or ""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("Дата составления:"):
            for run in list(paragraph.runs):
                run.text = ""
            run = paragraph.add_run(f"Дата составления: {data_akta}" if data_akta else "Дата составления: ____________________")
            run.font.name = "Times New Roman"
            break

    output_path = str(output_path)
    doc.save(output_path)
    return output_path
