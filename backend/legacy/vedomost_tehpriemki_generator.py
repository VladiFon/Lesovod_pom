# -*- coding: utf-8 -*-
"""Генератор батч-документа "Ведомость технической приёмки работ по
созданию лесных культур" (приложение 14 к постановлению Минлесхоза от
19.12.2016 N 80) — заполняет собственный шаблон
templates/vedomost_tehpriemki_shablon.xlsx на несколько участков сразу
(один лесхоз/год/сезон), по образцу akt_generator.py (копия шаблона +
openpyxl + ws.insert_rows под фактическое число участков).

Автозаполняются: №, квартал, выдел, площадь, ТЛУ+тип леса, категория
площади, главная порода/состав, количество посадочных мест на 1 га "по
данным технической приёмки" (из журнала участка). Остальные колонки
(способ обработки почвы, % отклонения, дефекты, оценка качества) —
без источника данных в базе, остаются пустыми для заполнения от руки
после печати (см. Этап 4 плана)."""
import os
import shutil

import openpyxl

_TEMPLATE_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "templates", "vedomost_tehpriemki_shablon.xlsx",
)

HEADER_ROW = 4
NUMBER_ROW = 5
DATA_START_ROW = 6


def generate_vedomost_tehpriemki(uchastki, lesnichestvo, god, sezon, yuridicheskoe_litso, output_path,
                                  template_path=None):
    """uchastki — список словарей из
    db.get_lesokultury_uchastki_for_tehpriemka() (не пуст)."""
    if not uchastki:
        raise ValueError("Нужен хотя бы один участок с записью «Техническая приёмка» в журнале")

    template_path = template_path or _TEMPLATE_PATH
    if not os.path.isfile(template_path):
        raise FileNotFoundError(f"Не найден шаблон ведомости техприёмки: {template_path}")

    output_path = str(output_path)
    shutil.copy(template_path, output_path)
    wb = openpyxl.load_workbook(output_path)
    ws = wb.active

    ws.cell(row=1, column=1).value = (
        f"ВЕДОМОСТЬ технической приёмки работ по созданию лесных культур, произведенных в {god or '____'} году"
    )
    ws.cell(row=2, column=2).value = sezon or ""
    ws.cell(row=2, column=5).value = lesnichestvo or ""
    ws.cell(row=2, column=11).value = yuridicheskoe_litso or ""

    n_extra_rows = len(uchastki) - 1
    if n_extra_rows > 0:
        ws.insert_rows(DATA_START_ROW + 1, amount=n_extra_rows)
    total_row = DATA_START_ROW + len(uchastki)

    for offset, u in enumerate(uchastki):
        row = DATA_START_ROW + offset
        tlu_tip = ", ".join(x for x in (u.get("tlu"), u.get("tip_lesa")) if x)
        poroda_sostav = ", ".join(x for x in (u.get("glavnaya_poroda"), u.get("sostav_formula")) if x)
        ws.cell(row=row, column=1, value=offset + 1)
        ws.cell(row=row, column=2, value=u.get("kvartal") or "")
        ws.cell(row=row, column=3, value=u.get("vydel") or "")
        ws.cell(row=row, column=4, value=u.get("ploshad") or "")
        ws.cell(row=row, column=6, value=tlu_tip)
        ws.cell(row=row, column=7, value=u.get("kategoriya_ploshadi") or "")
        ws.cell(row=row, column=10, value=poroda_sostav)
        ws.cell(row=row, column=17, value=u.get("kolichestvo_na_ga_tehpriemka"))

    ws.cell(row=total_row, column=1, value="Итого")
    ws.cell(row=total_row, column=4, value=sum(u.get("ploshad") or 0 for u in uchastki))

    wb.save(output_path)
    return output_path
