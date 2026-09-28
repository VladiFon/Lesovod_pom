# -*- coding: utf-8 -*-
"""Генератор "Полевой карточки инвентаризации участков лесных культур,
переводимых в покрытые лесом земли" (приложение 19 к постановлению
Минлесхоза от 19.12.2016 N 80) — заполняет заготовленный бланк
templates/kartochka_perevoda_shablon.docx данными участка/записи
журнала (см. lesokultury_field_card_fill.py), тем же способом, что и
остальные генераторы документов в проекте (шаблон + координаты ячеек)."""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from akt_template_fill import set_cell_text
import lesokultury_field_card_fill as fill
import lesokultury_normativy

_TEMPLATE_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "templates", "kartochka_perevoda_shablon.docx",
)


def generate_kartochka_perevoda(uchastok, meropriyatie, output_path, **manual_fields):
    """uchastok — dict из db.get_lesokultury_uchastok(); meropriyatie —
    dict из db.get_latest_lesokultury_perevod_meropriyatie() /
    get_lesokultury_meropriyatie() (обязательно с непустым dannye).

    manual_fields (все необязательны, по умолчанию ""):
        yuridicheskoe_litso, razmeshenie_v_ryadah_m, mezhdu_ryadami_m,
        zaklyuchenie_a, zaklyuchenie_b, zaklyuchenie_v, zaklyuchenie_g,
        zaklyuchenie_d, predsedatel: {dolzhnost, fio},
        chleny: [{dolzhnost, fio}, ...], data_akta.

    Бросает ValueError, если у meropriyatie нет структурированных данных
    полевой карточки (dannye) — печатать нечего."""
    dannye = (meropriyatie or {}).get("dannye")
    if not dannye:
        raise ValueError(
            "У выбранной записи журнала нет данных полевой карточки "
            "(проб/результатов обследования) — карточку перевода печатать не из чего."
        )

    template_path = manual_fields.pop("template_path", None) or _TEMPLATE_PATH
    if not os.path.isfile(template_path):
        raise FileNotFoundError(f"Не найден шаблон карточки перевода: {template_path}")

    doc = Document(template_path)
    tables = doc.tables
    if len(tables) < 6:
        raise ValueError("Шаблон карточки перевода повреждён — ожидались 6 таблиц")
    t_header, t_proby, t_rezultaty, t_normativ, t_zaklyuchenie, t_podpisi = tables[:6]

    proby = dannye.get("proby") or []
    rezultaty = dannye.get("rezultaty") or []
    itogo = dannye.get("itogo") or {}

    glavnaya_poroda = uchastok.get("glavnaya_poroda")
    tip_lesa = uchastok.get("tip_lesa")
    kolichestvo_na_ga = meropriyatie.get("kolichestvo_na_ga")
    srednyaya_vysota_m = meropriyatie.get("srednyaya_vysota_m")
    kolichestvo_tys_na_ga = kolichestvo_na_ga / 1000 if kolichestvo_na_ga is not None else None

    normativ_check = lesokultury_normativy.check_normativ_perevoda(
        glavnaya_poroda, tip_lesa, kolichestvo_tys_na_ga, srednyaya_vysota_m,
    )

    # --- Таблица 0: шапка (поля 1-10 приложения 19) ---
    set_cell_text(t_header, 0, 1, manual_fields.get("yuridicheskoe_litso", ""), align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t_header, 1, 1, uchastok.get("lesnichestvo") or "", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t_header, 2, 1, uchastok.get("kvartal") or "")
    set_cell_text(t_header, 2, 3, uchastok.get("vydel") or "")
    tlu_tip = ", ".join(x for x in (uchastok.get("tlu"), tip_lesa) if x)
    set_cell_text(t_header, 3, 1, tlu_tip, align=WD_ALIGN_PARAGRAPH.LEFT)
    vid_nasazhdeniya = uchastok.get("sostav_formula") or uchastok.get("glavnaya_poroda") or ""
    set_cell_text(t_header, 3, 3, vid_nasazhdeniya, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t_header, 4, 1, uchastok.get("god_sozdaniya") or "")
    set_cell_text(t_header, 4, 3, uchastok.get("ploshad") or "")
    shema = manual_fields.get("shema_smesheniya") or uchastok.get("sostav_formula") or ""
    ryad = manual_fields.get("razmeshenie_v_ryadah_m", "")
    mezhdu = manual_fields.get("mezhdu_ryadami_m", "")
    shema_text = f"{shema}; в рядах {ryad} м, между рядами {mezhdu} м" if (ryad or mezhdu) else shema
    set_cell_text(t_header, 5, 1, shema_text, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t_header, 6, 1, uchastok.get("gustota_posadki") or "", align=WD_ALIGN_PARAGRAPH.LEFT)

    # --- Таблицы 1-3: пробы / результаты / сверка с нормативом ---
    fill.fill_proba_table(t_proby, proby)
    fill.fill_rezultaty_table(t_rezultaty, rezultaty, itogo)
    fill.fill_normativ_table(
        t_normativ, normativ_check, glavnaya_poroda, tip_lesa,
        kolichestvo_tys_na_ga, srednyaya_vysota_m,
    )

    # --- Таблица 4: заключение комиссии (12 а-д, свободный ввод) ---
    fill.fill_zaklyuchenie_table(
        t_zaklyuchenie,
        a=manual_fields.get("zaklyuchenie_a", ""),
        b=manual_fields.get("zaklyuchenie_b", ""),
        v=manual_fields.get("zaklyuchenie_v", ""),
        g=manual_fields.get("zaklyuchenie_g", ""),
        d=manual_fields.get("zaklyuchenie_d", ""),
    )

    # --- Таблица 5: подписи комиссии ---
    fill.fill_signatures_table(t_podpisi, manual_fields.get("predsedatel"), manual_fields.get("chleny"))

    data_akta = manual_fields.get("data_akta") or meropriyatie.get("data") or ""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("Дата составления:"):
            for run in list(paragraph.runs):
                run.text = ""
            run = paragraph.add_run(f"Дата составления: {data_akta}" if data_akta else "Дата составления: ____________________")
            run.font.name = fill.FONT
            break

    output_path = str(output_path)
    doc.save(output_path)
    return output_path
