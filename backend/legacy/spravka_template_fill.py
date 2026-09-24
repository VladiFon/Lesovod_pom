# -*- coding: utf-8 -*-
"""
Заполнение "Справки о количестве заготовленной древесины" на основе
официального бланка (templates/spravka_shablon.docx) — по аналогии с
akt_template_fill.py: текст точечно вписывается в заранее определённые
ячейки трёх таблиц бланка (шапка / таблица объёмов / подпись), всё
остальное (рамки, жирные подписи, объединения ячеек) остаётся как в
оригинале, поэтому внешний вид документа не меняется от заполнения.

В отличие от прежней версии (построение документа с нуля, произвольный
набор колонок-пород) бланк содержит ФИКСИРОВАННЫЙ набор колонок по
породам: Ель, Сосна, Береза, Осина, Дуб, Ольха С[ерая], Липа — плюс
"Всего". Порода, не входящая в этот список, попадает только в "Всего" —
отдельной колонки под неё в официальном бланке нет.
"""
from akt_template_fill import set_cell_text
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

# код породы (см. PORODA_NAMES в spravka_generator.py) -> колонка таблицы 1
PORODA_COL = {
    "Е": 4, "С": 5, "Б": 6, "ОС": 7, "Д": 8,
    "Олс": 9, "ОЛС": 9,
    "Лп": 10,
}


def fill_spravka(template_path, data, output_path):
    doc = Document(template_path)
    t0, t1, t2 = doc.tables[0], doc.tables[1], doc.tables[2]
    d = data

    # --- шапка ---
    set_cell_text(t0, 0, 2, d.get("lesxoz", ""), align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t0, 2, 1, d.get("lesopolzovatel", ""), align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t0, 5, 7, d.get("bilet_nomer", ""))
    set_cell_text(t0, 5, 12, f"«{d.get('bilet_day','')}»" if d.get("bilet_day") else "")
    set_cell_text(t0, 5, 13, d.get("bilet_month", ""))
    set_cell_text(t0, 5, 14, (d.get("bilet_year", "") + " г.") if d.get("bilet_year") else "",
                  align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t0, 6, 4, d.get("kvartal", ""))
    set_cell_text(t0, 6, 9, d.get("vydel", ""))
    set_cell_text(t0, 6, 11, d.get("lesnichestvo", ""), align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t0, 8, 4, d.get("delyanka_nomer", ""))

    # --- таблица объёмов (значения по фиксированным породам) ---
    def fill_row(row, total, by_poroda=None):
        set_cell_text(t1, row, 3, total)
        if by_poroda:
            written_cols = set()
            for code, value in by_poroda.items():
                col = PORODA_COL.get(code)
                if col is not None and col not in written_cols:
                    set_cell_text(t1, row, col, value)
                    written_cols.add(col)

    fill_row(3, d.get("ploshad_proydennaya", ""))
    fill_row(4, d.get("itogo_total", ""), d.get("itogo_by_poroda"))
    fill_row(5, d.get("delovaya_total", ""), d.get("delovaya_by_poroda"))
    fill_row(6, d.get("kr_total", ""), d.get("kr_by_poroda"))
    fill_row(7, d.get("sr_total", ""), d.get("sr_by_poroda"))
    fill_row(8, d.get("ml_total", ""), d.get("ml_by_poroda"))
    fill_row(9, d.get("drova_total", ""), d.get("drova_by_poroda"))
    fill_row(10, d.get("likvid_such_krony", ""))
    fill_row(11, d.get("hvorost_total", ""), d.get("hvorost_by_poroda"))
    fill_row(12, d.get("vsego_total", ""), d.get("vsego_by_poroda"))

    # --- подпись ---
    set_cell_text(t2, 1, 4, d.get("rukovoditel_fio", ""))

    output_path = str(output_path)
    doc.save(output_path)
    return output_path
