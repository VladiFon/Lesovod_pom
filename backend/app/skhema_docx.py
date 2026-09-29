# -*- coding: utf-8 -*-
"""
Лист «5.10. Схема разработки лесосеки» строго по шаблону
(legacy/templates/skhema_5_10_shablon.docx — «Шаблон схема-чертежа.docx»):
заголовок и «Условные обозначения» берутся из самого шаблона как есть,
в пустое поле между ними встаёт чертёж из абриса (abris_tool.html
сохраняет только поле чертежа — DRAW_AREA_MM, 177,5 × 200 мм).

Используется двумя путями:
  - build_docx() — отдельный Word-файл схемы (кнопка в абрисе);
  - insert_page() — тот же лист внутри техкарты (tehkarta_generator).
"""
import copy
import io
import json
from pathlib import Path
from typing import Optional

TEMPLATE_NAME = "skhema_5_10_shablon.docx"
# Размер поля чертежа — как DRAW_AREA_MM в abris_tool.html.
DRAW_W_CM = 17.75
DRAW_H_CM = 20.0

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"


def template_path(templates_dir: Optional[str] = None) -> Path:
    """Шаблон ищем в каталоге бланков, иначе — рядом с кодом (legacy/templates)."""
    candidates = []
    if templates_dir:
        candidates.append(Path(templates_dir) / TEMPLATE_NAME)
    candidates.append(Path(__file__).resolve().parent.parent / "legacy" / "templates" / TEMPLATE_NAME)
    for p in candidates:
        if p.exists():
            return p
    raise FileNotFoundError(f"Не найден шаблон схемы {TEMPLATE_NAME}")


def is_template_drawing(proekt_json: Optional[str]) -> bool:
    """Картинка абриса — только поле чертежа (новый абрис), а не целый лист А4."""
    if not proekt_json:
        return False
    try:
        return int(json.loads(proekt_json).get("version", 0)) >= 2
    except (ValueError, TypeError, AttributeError):
        return False


def _legend_start(doc):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("Условные обозначения"):
            return i
    raise ValueError("В шаблоне схемы не найден блок «Условные обозначения»")


def _png_size(image_path: str):
    """(ширина, высота) PNG из заголовка — без Pillow (на сервере его может не быть)."""
    import struct

    with open(image_path, "rb") as f:
        head = f.read(24)
    if head[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError("Картинка абриса — не PNG")
    return struct.unpack(">II", head[16:24])


def _fit(image_path: str, max_w_cm: float, max_h_cm: float):
    from docx.shared import Cm

    w, h = _png_size(image_path)
    k = min(max_w_cm / w, max_h_cm / h)
    return Cm(w * k), Cm(h * k)


def build_docx(image_path: str, templates_dir: Optional[str] = None) -> bytes:
    """Word-файл схемы: шаблон, в котором пустое поле заменено чертежом."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document(str(template_path(templates_dir)))
    start = _legend_start(doc)
    paras = doc.paragraphs
    # абзацы 1..start-1 — пустые строки под чертёж: оставляем один под картинку
    for p in paras[2:start]:
        p._element.getparent().remove(p._element)
    img_p = paras[1]
    for r in list(img_p.runs):
        r._element.getparent().remove(r._element)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    w, h = _fit(image_path, DRAW_W_CM, DRAW_H_CM)
    img_p.add_run().add_picture(image_path, width=w, height=h)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _legend_elements(templates_dir: Optional[str] = None):
    """Копии XML-абзацев «Условные обозначения» и ниже (со знаками-фигурами)."""
    from docx import Document

    src = Document(str(template_path(templates_dir)))
    start = _legend_start(src)
    return [copy.deepcopy(p._element) for p in src.paragraphs[start:]]


def insert_page(doc, before_paragraph, image_path: str, templates_dir: Optional[str] = None,
                max_h_cm: Optional[float] = None):
    """Вставляет перед before_paragraph лист как в шаблоне: заголовок,
    чертёж, условные обозначения (копия из шаблона)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    title = before_paragraph.insert_paragraph_before()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("5.10. Схема разработки лесосеки")
    run.bold = True

    img_p = before_paragraph.insert_paragraph_before()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    w, h = _fit(image_path, DRAW_W_CM, max_h_cm or DRAW_H_CM)
    img_p.add_run().add_picture(image_path, width=w, height=h)

    # id фигур должны быть уникальны в документе
    used = {int(e.get("id")) for e in doc.element.body.iter(_WP + "docPr") if (e.get("id") or "").isdigit()}
    next_id = max(used | {0}) + 1000
    anchor = before_paragraph._element
    for el in _legend_elements(templates_dir):
        for pr in el.iter(_WP + "docPr"):
            pr.set("id", str(next_id))
            next_id += 1
        anchor.addprevious(el)
