# -*- coding: utf-8 -*-
"""Сводки экрана «Расход / ЕГАИС» (30.09.2026, запрос лесничего):

1. Расход по ЕГАИС — кто оформил расход, на какой объём, кому (грузополучатель),
   какую продукцию и с какой делянки.
2. Приход по ЕГАИС — сколько прихода, какой продукции, с какой делянки.
3. Приход по нарядам — то же по нарядам-заданиям, введённым в приложении.

Во всех трёх итоги делятся на деловую древесину, дрова и общий итог.

Делянка строки ЕГАИС определяется ТОЛЬКО через привязку склада
(raskhod_v2.resolve_egais_sklady), как и во всём остальном учёте ЕГАИС, —
не по кварталу/выделу строки. Строки складов без однозначной привязки
показываются с пометкой «склад не привязан», а не приписываются наугад.

Бэкенд отдаёт строки, уже схлопнутые до «документ × продукция» (брёвна
поштучного учёта одного документа — в одну строку), и списки значений для
фильтров; фильтрация по делянке/сотруднику/получателю и группировка
делаются на клиенте — объём данных за год (≈10 тыс. строк выгрузки,
≈2 тыс. после схлопывания) это позволяет без задержек.
"""
import re
from io import BytesIO

from app import legacy_bridge  # noqa: F401
import raskhod_v2 as legacy_raskhod

GROUP_DELOVAYA = "delovaya"
GROUP_DROVA = "drova"
GROUP_HVOROST = "hvorost"

GROUP_LABELS = {
    GROUP_DELOVAYA: "Деловая древесина",
    GROUP_DROVA: "Дрова",
    GROUP_HVOROST: "Хворост",
}

SORTIMENT_LABELS = {"KR": "Крупная", "SR": "Средняя", "ML": "Мелкая", "DROVA": "Дрова", "HVOROST": "Хворост"}


def _date_iso(value):
    """'ДД.ММ.ГГГГ' / 'ГГГГ-ММ-ДД' -> 'ГГГГ-ММ-ДД' (пусто, если не распознано)."""
    text = str(value or "").strip()[:10]
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return text
    return legacy_raskhod._egais_date_sort_key(text)


def _in_period(date_iso, date_from, date_to):
    if not date_from and not date_to:
        return True
    if not date_iso:
        return False
    if date_from and date_iso < date_from:
        return False
    if date_to and date_iso > date_to:
        return False
    return True


def egais_group(godnost, nomenklatura):
    text = f"{godnost or ''} {nomenklatura or ''}".lower()
    if "дров" in text:
        return GROUP_DROVA
    return GROUP_DELOVAYA


def egais_produkciya(nomenklatura, poroda="", sort=""):
    """«Лесоматериалы круглые шт, Ель, 6 м., 24 см, D» -> «Лесоматериалы
    круглые, Ель, 6 м., D»: диаметр отдельного бревна для сводки не нужен
    (иначе каждая строка — своя «продукция»), а «шт»/«гр» — способ учёта,
    а не продукция."""
    parts = [p.strip() for p in str(nomenklatura or "").split(",") if p.strip()]
    if not parts:
        return ", ".join(x for x in (poroda, sort) if x) or "—"
    parts[0] = re.sub(r"\s+(шт|гр)\.?$", "", parts[0])
    parts = [p for p in parts if not re.search(r"\d\s*см\b", p)]
    return ", ".join(parts)


def _item_labels(conn):
    items = legacy_raskhod._load_items_for_sklad_match(conn)
    return {
        it["id"]: {
            "item_id": it["id"],
            "delyanka_id": it["delyanka_id"],
            "delyanka": it.get("delyanka_nazvanie") or f"Делянка №{it['delyanka_id']}",
            "vydel_label": legacy_raskhod.egais_item_label(it),
        }
        for it in items
    }


def egais_svodka(conn, napravlenie="raskhod", date_from=None, date_to=None,
                 bez_povtornogo_prihoda=True):
    """napravlenie: 'raskhod' — все «Расход …» документы, 'prihod' — «Приход».
    bez_povtornogo_prihoda: не считать приход на склад-получатель при
    внутреннем перемещении (это та же древесина, уже учтённая приходом на
    складе делянки — см. raskhod_v2._egais_linked_doc_numbers)."""
    resolution = legacy_raskhod.resolve_egais_sklady(conn)
    labels = _item_labels(conn)
    if napravlenie == "prihod":
        where = "tip_dokumenta = 'Приход'"
    else:
        where = "tip_dokumenta LIKE 'Расход%'"
    linked = legacy_raskhod._egais_linked_doc_numbers(conn) if (
        napravlenie == "prihod" and bez_povtornogo_prihoda) else set()

    rows = conn.execute(
        "SELECT data_dokumenta, data_dokumenta_sort, tip_dokumenta, nomer_dokumenta, "
        "kvartal, vydel, sklad, sklad_kontragent, COALESCE(gruzopoluchatel, ''), poroda, sort, "
        "tehnicheskaya_godnost, nomenklatura, obyom, osnovanie, nomer_osnovaniya, sotrudnik "
        f"FROM egais_operation WHERE {where}"
    ).fetchall()

    agg = {}
    skipped_povtor = 0
    for (data, data_sort, tip, nomer, kvartal, vydel, sklad, kontragent, poluchatel, poroda, sort,
         godnost, nomenklatura, obyom, osnovanie, nomer_osn, sotrudnik) in rows:
        date_iso = data_sort or _date_iso(data)
        if not _in_period(date_iso, date_from, date_to):
            continue
        if linked and nomer in linked:
            skipped_povtor += 1
            continue
        key_sklad = legacy_raskhod.egais_sklad_key(sklad, kvartal, vydel)
        info = resolution.get(key_sklad) or {}
        item_id = info.get("item_id")
        lab = labels.get(item_id) if item_id else None
        if not poluchatel and napravlenie != "prihod":
            poluchatel = kontragent or ""
        osn = " ".join(x for x in (osnovanie or "", nomer_osn or "") if x).strip()
        produkciya = egais_produkciya(nomenklatura, poroda, sort)
        group = egais_group(godnost, nomenklatura)
        k = (date_iso, nomer or "", tip or "", sotrudnik or "", poluchatel, key_sklad,
             poroda or "", produkciya, group, osn)
        row = agg.get(k)
        if row is None:
            row = agg[k] = {
                "data": data or "",
                "data_iso": date_iso,
                "nomer": nomer or "",
                "tip": tip or "",
                "sotrudnik": sotrudnik or "",
                "poluchatel": poluchatel,
                "osnovanie": osn,
                "sklad": key_sklad,
                "kvartal": kvartal or "",
                "vydel": vydel or "",
                "item_id": item_id,
                "delyanka_id": lab["delyanka_id"] if lab else None,
                "delyanka": lab["delyanka"] if lab else "",
                "vydel_label": lab["vydel_label"] if lab else "",
                "poroda": poroda or "",
                "produkciya": produkciya,
                "group": group,
                "obyom": 0.0,
            }
        row["obyom"] += abs(obyom or 0.0)

    out = sorted(agg.values(), key=lambda r: (r["data_iso"], r["nomer"], r["produkciya"]), reverse=True)
    for r in out:
        r["obyom"] = round(r["obyom"], 3)
    return {"rows": out, "skipped_povtor": skipped_povtor}


def naryady_svodka(conn, date_from=None, date_to=None):
    labels = _item_labels(conn)
    rows = conn.execute(
        "SELECT n.id, n.data, n.nomer_naryada, n.item_id, n.delyanka_id, n.primechanie, "
        "p.poroda, p.sortiment, p.obyom, di.kvartal, di.vydel "
        "FROM raskhod_pozitsiya p JOIN raskhod_naryad n ON n.id = p.naryad_id "
        "LEFT JOIN delyanka_item di ON di.id = n.item_id"
    ).fetchall()
    out = []
    for (nid, data, nomer, item_id, delyanka_id, primechanie, poroda, sortiment, obyom,
         kvartal, vydel) in rows:
        date_iso = _date_iso(data)
        if not _in_period(date_iso, date_from, date_to):
            continue
        lab = labels.get(item_id) or {}
        sortiment = (sortiment or "").upper()
        group = GROUP_DROVA if sortiment == "DROVA" else GROUP_HVOROST if sortiment == "HVOROST" else GROUP_DELOVAYA
        poroda_label = legacy_raskhod._poroda_label(poroda) if poroda else ""
        out.append({
            "naryad_id": nid,
            "data": data or "",
            "data_iso": date_iso,
            "nomer": nomer or "",
            "primechanie": primechanie or "",
            "item_id": item_id,
            "delyanka_id": lab.get("delyanka_id", delyanka_id),
            "delyanka": lab.get("delyanka") or (f"Делянка №{delyanka_id}" if delyanka_id else ""),
            "vydel_label": lab.get("vydel_label") or f"кв.{kvartal or '—'} выд.{vydel or '—'}",
            "poroda": poroda_label or (poroda or ""),
            "sortiment": sortiment,
            "produkciya": f"{poroda_label or poroda or '—'}, {SORTIMENT_LABELS.get(sortiment, sortiment or '—').lower()}",
            "group": group,
            "obyom": round(float(obyom or 0.0), 3),
        })
    out.sort(key=lambda r: (r["data_iso"], r["nomer"]), reverse=True)
    return {"rows": out}


def svodka_to_xlsx(title, subtitle, columns, sections):
    """Excel ровно с тем, что показано на экране (фильтры/группировка уже
    применены на клиенте). columns: [{"key", "label"}]; sections: [{"title",
    "rows": [{key: value}], "total_label", "total"}] + последняя — «Итого»."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    if subtitle:
        ws.append([subtitle])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append([c.get("label", "") for c in columns])
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EFE9")
        cell.border = border
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    def put(values, bold=False, fill=None):
        ws.append(values)
        for cell in ws[ws.max_row]:
            cell.border = border
            if bold:
                cell.font = Font(bold=True)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
            if isinstance(cell.value, float):
                cell.number_format = "0.000"

    ncol = len(columns)
    for sec in sections:
        if sec.get("title"):
            put([sec["title"]] + [""] * (ncol - 1), bold=True, fill="F4F1E8")
        for r in sec.get("rows") or []:
            put([r.get(c["key"], "") for c in columns])
        if sec.get("total_label") is not None:
            put([sec["total_label"]] + [""] * (ncol - 2) + [sec.get("total", 0.0)], bold=True, fill="EFEFEF")

    for i, c in enumerate(columns, start=1):
        width = c.get("width") or (12 if c["key"] == "obyom" else 22)
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = width
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
