# -*- coding: utf-8 -*-
"""Загрузка «Книги производства лесных культур» (.xls/.xlsx, лист на каждый
год посадки) в «Лесные культуры».

Три шага (роутер app/routers/lesokultury.py):
1. parse_book() — разбор листов выбранных лет в строки-участки. Колонки у
   разных лет сдвинуты, поэтому столбцы ищутся по тексту шапки (строки 3–4),
   а не по жёстким номерам. В листах 2017–2019 место участка записано тремя
   колонками: квартал, старый выдел, новый выдел по действующей таксации
   (в 2017 новый бывает в одной из двух соседних колонок).
2. plan_import() — сверка с базой без записи: какие участки новые, какие уже
   есть (по метке источника, по метке старого скрипта
   import_lesokultury_xls.py или по кварталу+выделам+году), что спорно.
3. apply_import() — запись: новые участки создаются, у найденных
   дописываются только пустые поля, в журнал добавляются инвентаризации,
   перевод, доращивание, списание из книги (без повторов).
"""
import io
import json
import re
from typing import Dict, Iterable, List, Optional

from app import legacy_bridge  # noqa: F401
import db as legacy_db

TIP_INV_1 = "Инвентаризация 1-го года"
TIP_INV_3 = "Инвентаризация 3-го года"
TIP_PEREVOD = "Перевод в покрытые лесом земли"
TIP_DORASHCHIVANIE = "Доращивание"
TIP_SPISANIE = "Списание"

STATUS_ACTIVE = "активен"
STATUS_PEREVEDEN = "переведён"
STATUS_SPISAN = "списан"

_SOSTAV_RE = re.compile(r"(?i)^\d{1,2}\s*[А-ЯЁ][а-яё]?(\s*\d{0,2}\s*[А-ЯЁ][а-яё]?|\+[А-ЯЁ][а-яё]?)*$")
_SOSTAV_IN_TEXT_RE = re.compile(r"(?i)\d{1,2}\s*[а-яё]{1,2}(?:\s*\d{0,2}\s*\+?\s*[а-яё]{1,2})*")
_TLU_RE = re.compile(r"(?i)(^|[^а-я])[А-Д]\s*\d|кис|чер|бр|сл|мш")
_SCHEME_RE = re.compile(r"^(\d+[.,]?\d*)\s*[*xх×]\s*(\d+[.,]?\d*)$")
_YEAR_RE = re.compile(r"20\d\d")
_SEASONS = ("весна", "осень", "лето", "зима")


class KnigaError(Exception):
    pass


# --------------------------------------------------------------------------- #
#   Чтение файла
# --------------------------------------------------------------------------- #
class _Sheet:
    def __init__(self, name: str, rows: List[list]):
        self.name = name
        self.rows = rows

    @property
    def nrows(self) -> int:
        return len(self.rows)

    def row_values(self, r: int) -> list:
        return self.rows[r]


def open_book(content: bytes, filename: str) -> List[_Sheet]:
    name = (filename or "").lower()
    try:
        if name.endswith(".xlsx") or content[:2] == b"PK":
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            sheets = []
            for ws in wb.worksheets:
                rows = [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
                width = max((len(r) for r in rows), default=0)
                sheets.append(_Sheet(ws.title, [r + [""] * (width - len(r)) for r in rows]))
            return sheets
        import xlrd

        wb = xlrd.open_workbook(file_contents=content)
        return [_Sheet(s.name, [s.row_values(r) for r in range(s.nrows)]) for s in wb.sheets()]
    except ImportError as exc:
        raise KnigaError(f"На сервере не установлен модуль для чтения Excel ({exc.name}): pip install -r requirements.txt")
    except Exception as exc:  # noqa: BLE001 — битый/не тот файл
        raise KnigaError(f"Не удалось открыть файл как книгу Excel: {exc}")


# --------------------------------------------------------------------------- #
#   Разбор
# --------------------------------------------------------------------------- #
def _s(v) -> str:
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else repr(round(v, 4))
    return str(v).strip()


def _num(v) -> Optional[float]:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ".").strip())
    except ValueError:
        return None


def _find(header: List[str], key: str) -> Optional[int]:
    for i, x in enumerate(header):
        if key in x:
            return i
    return None


def vydel_tokens(text: str) -> List[str]:
    """«19,20,23» / «2.3.10.» / «12.13» -> ["19","20","23"] … — выделы в книге
    перечисляют и через запятую, и через точку (Excel делает из «12.13»
    число)."""
    return [t for t in re.split(r"[\s,.;+]+", text or "") if t]


def _sheet_year(name: str) -> Optional[int]:
    m = re.match(r"\s*(20\d\d)", name or "")
    return int(m.group(1)) if m else None


def _parse_sheet(sh: _Sheet, year: int) -> List[dict]:
    if sh.nrows < 7:
        return []
    h2 = [_s(x).lower() for x in sh.row_values(2)]
    h3 = [_s(x).lower() for x in sh.row_values(3)]
    c_god = _find(h2, "год и сезон")
    c_mesto = _find(h2, "мест")
    c_pl = _find(h2, "площадь")
    c_obr = _find(h2, "обработки")
    c_por = _find(h2, "главная порода")
    c_prin = _find(h2, "принято")
    c_spis = _find(h2, "списано")
    c_per = _find(h2, "переведено")
    c_uh = _find(h2, "уходов")
    c_posev = _find(h3, "посев")
    c_posad = _find(h3, "посадка")
    if None in (c_mesto, c_pl, c_por, c_prin, c_spis, c_per):
        raise KnigaError(f"Лист «{sh.name}»: не нашёл шапку книги л/к (Местонахождение, Площадь, Главная порода…)")

    out: List[dict] = []
    last = None
    for r in range(6, sh.nrows):
        row = list(sh.row_values(r))
        mesto = [_s(x) for x in row[c_mesto:c_pl]]
        if not mesto:
            continue
        kv = mesto[0]
        if not re.match(r"^\d+$", kv):
            continue
        pl = _num(row[c_pl])

        tlu = ""
        vyd_old = mesto[1] if len(mesto) > 1 else ""
        vyd_new: List[str] = []
        note = ""
        if _TLU_RE.search(vyd_old):
            tlu, vyd_old = vyd_old, ""
        for x in mesto[2:]:
            if not x:
                continue
            if _TLU_RE.search(x) and not tlu:
                tlu = x
            elif re.search(r"[А-Яа-я]", x):
                m2 = re.match(r"^([\d.,]+)\s+(.*)$", x)
                if m2:
                    vyd_new.append(m2.group(1))
                    note = m2.group(2)
                else:
                    note = x
            else:
                vyd_new.append(x)
        vyd_new_text = ",".join(vyd_new)

        # строка-продолжение (выделы перенесены на следующую строку без площади)
        if pl is None and last is not None and last["kvartal"] == kv:
            last["vydel_staryy_kniga"] = ", ".join(v for v in (last["vydel_staryy_kniga"], vyd_old) if v)
            if vyd_new_text:
                last["vydel_novyy_kniga"] = ", ".join(v for v in (last["vydel_novyy_kniga"], vyd_new_text) if v)
            continue
        if pl is None or pl <= 0:
            continue

        rec = {
            "key": f"{year}/стр.{r + 1}",
            "list": sh.name.strip().rstrip(". ").strip(),
            "god": year,
            "stroka": r + 1,
            "god_sezon": _s(row[c_god]) if c_god is not None else "",
            "kvartal": kv,
            "vydel_staryy_kniga": vyd_old,
            "vydel_novyy_kniga": vyd_new_text,
            "vydel_primechanie": note,
            "tlu": tlu,
            "ploshad": round(pl, 2),
            "obrabotka": _s(row[c_obr]) if c_obr is not None else "",
            "metod": ("посев" if c_posev is not None and _num(row[c_posev]) else "")
            or ("посадка" if c_posad is not None and _num(row[c_posad]) else ""),
            "glavnaya_poroda": _s(row[c_por]),
            "problemy": [],
        }

        sostav = mezhdu = v_ryadu = material = ""
        gustota = None
        prizh: List[float] = []
        for x in row[c_por + 1:c_prin]:
            t = _s(x)
            if not t:
                continue
            n = _num(x)
            mm = _SCHEME_RE.match(t)
            if mm:
                mezhdu, v_ryadu = mm.group(1).replace(",", "."), mm.group(2).replace(",", ".")
                continue
            if not sostav and _SOSTAV_RE.match(t.replace(" ", "")):
                sostav = t.replace(" ", "")
                continue
            if n is not None:
                if n >= 1000:
                    gustota = int(n)
                elif not mezhdu and n < 6:
                    mezhdu = t
                elif mezhdu and not v_ryadu and n < 3 and gustota is None:
                    v_ryadu = t
                elif n <= 100:
                    prizh.append(n)
                continue
            material = f"{material} {t}".strip()
        if not sostav:
            mm = re.search(r"\d{1,2}[А-ЯЁ][а-яё]?(?:\d{1,2}[А-ЯЁ][а-яё]?)+", material)
            if mm:
                sostav = mm.group(0)
        rec.update(
            sostav=sostav,
            shema_mezhdu_ryadami=_num(mezhdu) if mezhdu else None,
            shema_v_ryadu=_num(v_ryadu) if v_ryadu else None,
            gustota=gustota,
            material=material,
            prizhivaemost_1=prizh[0] if prizh else None,
            prizhivaemost_2=prizh[1] if len(prizh) > 1 else None,
        )

        spis = [_s(x) for x in row[c_spis:c_per] if _s(x)]
        rec["spisano_kniga"] = " | ".join(spis)
        per_end = c_uh if c_uh is not None and c_uh > c_per else c_per + 2
        per = [_s(x) for x in row[c_per:per_end] if _s(x)]
        rec["perevod_kniga"] = " | ".join(per)

        p_area = None
        p_card = p_sostav = p_year = ""
        dorashch = spisat = ""
        for t in per:
            low = t.lower()
            if "доращ" in low:
                dorashch = t
                continue
            if "списа" in low:
                spisat = t
                continue
            mm = re.search(r"№\s*(\d+)", t)
            ys = _YEAR_RE.findall(t)
            if mm or ys:
                p_card = mm.group(1) if mm else ""
                p_year = ys[-1] if ys else ""
                rest = re.sub(r"№\s*\d+|20\d\d", "", t)
                ms = _SOSTAV_IN_TEXT_RE.search(rest)
                if ms:
                    p_sostav = ms.group(0).replace(" ", "").upper().replace("ОС", "Ос").replace("ЛП", "Лп")
            elif _num(t) is not None and p_area is None:
                p_area = _num(t)
        if p_area is not None and not p_year and not dorashch and not spisat:
            rec["problemy"].append("в колонке перевода есть площадь, но нет года и № карточки")
        spisano_ranee = bool(rec["spisano_kniga"]) and "огораж" not in rec["spisano_kniga"].lower()
        if spisano_ranee and dorashch:
            rec["problemy"].append("в книге и списание, и доращивание")
        rec.update(
            perevod_god=p_year,
            perevod_kartochka=p_card,
            perevod_sostav=p_sostav,
            perevod_ploshad=(p_area if p_year else None) or (rec["ploshad"] if p_year else None),
            dorashchivanie=dorashch,
            spisat=spisat,
            spisano_ranee=spisano_ranee,
        )
        if spisano_ranee:
            status = STATUS_SPISAN
        elif spisat:
            status = STATUS_SPISAN
        elif dorashch:
            status = STATUS_ACTIVE
        elif p_year:
            status = STATUS_PEREVEDEN
        else:
            status = STATUS_ACTIVE
        rec["status"] = status
        if not sostav:
            rec["problemy"].append("не найден состав")
        out.append(rec)
        last = rec
    return out


def parse_book(sheets: List[_Sheet], god_from: int = 2017, god_to: int = 2100) -> List[dict]:
    rows: List[dict] = []
    for sh in sheets:
        year = _sheet_year(sh.name)
        if year is None or year < god_from or year > god_to:
            continue
        rows.extend(_parse_sheet(sh, year))
    if not rows:
        raise KnigaError(f"В книге нет участков за {god_from}–{god_to} (листы называются годом посадки: 2017, 2018 …)")
    return rows


# --------------------------------------------------------------------------- #
#   Сверка с базой
# --------------------------------------------------------------------------- #
def _norm_name(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def _taxation_vydely(conn, lesnichestvo: str) -> Optional[Dict[int, set]]:
    """{квартал: {выделы}} таксации лесничества или None, если её нет."""
    target = _norm_name(lesnichestvo)
    try:
        rows = conn.execute(
            "SELECT l.name, k.nomer, v.nomer FROM vydel v "
            "JOIN kvartal k ON k.id = v.kvartal_id JOIN lesnichestvo l ON l.id = k.lesnichestvo_id"
        ).fetchall()
    except Exception:  # noqa: BLE001 — таблиц таксации может не быть
        return None
    result: Dict[int, set] = {}
    for name, kv, vd in rows:
        n = _norm_name(name)
        if n != target and target not in n and n not in target:
            continue
        result.setdefault(int(kv), set()).add(str(vd).strip())
    return result or None


def _resolve_vydel(kv: str, raw: str, taxation: Optional[Dict[int, set]]) -> (str, List[str]):
    """Выдел для участка: «12, 13». Если в таксации есть ровно «12.13»
    (подвыдел), он и остаётся. Возвращает (текст, выделы, которых нет в
    таксации)."""
    raw = (raw or "").strip()
    if not raw:
        return "", []
    known = taxation.get(int(kv), set()) if taxation is not None else None
    if known is not None and raw.rstrip(".") in known:
        return raw.rstrip("."), []
    tokens = vydel_tokens(raw)
    missing = [t for t in tokens if known is not None and t not in known]
    return ", ".join(tokens), missing


def _uchastki(conn) -> List[dict]:
    return legacy_db.get_lesokultury_uchastki(conn, include_spisannye=True)


def plan_import(conn, rows: List[dict], lesnichestvo: str, overrides: Optional[Dict[str, str]] = None) -> dict:
    """Для каждой строки книги: действие («новый» / «есть в базе»), id
    найденного участка, итоговый выдел и проблемы. В базу не пишет."""
    overrides = overrides or {}
    taxation = _taxation_vydely(conn, lesnichestvo)
    existing = _uchastki(conn)
    by_istochnik = {u.get("istochnik"): u for u in existing if u.get("istochnik")}
    lesn = _norm_name(lesnichestvo)
    used: set = set()

    for rec in rows:
        raw = overrides.get(rec["key"]) or rec["vydel_novyy_kniga"] or rec["vydel_staryy_kniga"]
        vydel, missing = _resolve_vydel(rec["kvartal"], raw, taxation)
        rec["vydel"] = vydel
        rec["vydel_staryy"] = (
            ", ".join(vydel_tokens(rec["vydel_staryy_kniga"])) if rec["vydel_novyy_kniga"] else ""
        )
        problems = [p for p in rec["problemy"]]
        if missing:
            problems.append(f"в таксации кв. {rec['kvartal']} нет выдела {', '.join(missing)}")
        if re.match(r"^\d+\.\d+\.?$", raw.strip()) and rec["key"] not in overrides:
            problems.append(f"выдел «{raw}» записан через точку — понял как {vydel}")
        rec["problemy_vse"] = problems

        match = by_istochnik.get(f"книга л/к: {rec['key']}")
        how = "загружен ранее"
        if match is None:
            tag = f"[импорт xls: {rec['god']}/стр.{rec['stroka']}]"
            match = next((u for u in existing if tag in (u.get("primechaniya") or "")), None)
            how = "загружен старым скриптом"
        if match is None:
            want = {frozenset(vydel_tokens(vydel)), frozenset(vydel_tokens(rec["vydel_staryy_kniga"]))}
            for u in existing:
                if u["id"] in used or (lesn and _norm_name(u.get("lesnichestvo")) not in (lesn, "")):
                    continue
                if str(u.get("kvartal") or "").strip() != rec["kvartal"]:
                    continue
                if str(rec["god"]) not in str(u.get("god_sozdaniya") or ""):
                    continue
                if frozenset(vydel_tokens(str(u.get("vydel") or ""))) in want:
                    match = u
                    how = "найден по кварталу, выделу и году"
                    break
        if match is not None:
            used.add(match["id"])
            rec["deystvie"] = "есть в базе"
            rec["uchastok_id"] = match["id"]
            rec["kak_nayden"] = how
        else:
            rec["deystvie"] = "новый"
            rec["uchastok_id"] = None
            rec["kak_nayden"] = ""

    summary = {
        "vsego": len(rows),
        "novyh": sum(1 for r in rows if r["deystvie"] == "новый"),
        "est_v_baze": sum(1 for r in rows if r["deystvie"] == "есть в базе"),
        "s_problemami": sum(1 for r in rows if r["problemy_vse"]),
        "ploshad": round(sum(r["ploshad"] for r in rows), 1),
        "perevod": _sum_status(rows, lambda r: r["status"] == STATUS_PEREVEDEN),
        "dorashchivanie": _sum_status(rows, lambda r: bool(r["dorashchivanie"]) and r["status"] == STATUS_ACTIVE),
        "spisanie": _sum_status(rows, lambda r: r["status"] == STATUS_SPISAN),
        "taksatsiya_proverena": taxation is not None,
    }
    return {"summary": summary, "rows": rows}


def _sum_status(rows, pred) -> dict:
    sel = [r for r in rows if pred(r)]
    return {"uchastkov": len(sel), "ploshad": round(sum(r["ploshad"] for r in sel), 1)}


# --------------------------------------------------------------------------- #
#   Запись
# --------------------------------------------------------------------------- #
def _metod_sozdaniya(rec: dict) -> str:
    material = rec["material"].lower()
    if rec["metod"] == "посадка" and "зкс" in material:
        return "созданы ЗКС"
    return rec["metod"]


def _sposob_obrabotki(rec: dict) -> str:
    text = rec["obrabotka"].strip()
    # «мех» в книге — механизированная обработка бороздами (уточнено 28.09.2026)
    if text.lower().startswith("мех"):
        return "бороздами"
    return text


def uchastok_fields(rec: dict, lesnichestvo: str) -> dict:
    prim = [f"Загружено из книги л/к (лист {rec['list']}, стр. {rec['stroka']})."]
    god_sezon = rec["god_sezon"]
    if god_sezon and not re.match(r"^20\d\d$", god_sezon):
        if god_sezon.lower() in _SEASONS:
            prim.append(f"Сезон: {god_sezon.lower()}.")
        else:
            prim.append(f"Исходные данные: {god_sezon}.")
    if rec["vydel_primechanie"]:
        prim.append(f"К выделу: {rec['vydel_primechanie']}.")
    if rec["spisano_kniga"] and not rec["spisano_ranee"]:
        prim.append(rec["spisano_kniga"] + ".")
    return {
        "lesnichestvo": lesnichestvo,
        "kvartal": rec["kvartal"],
        "vydel": rec["vydel"],
        "vydel_staryy": rec["vydel_staryy"] or None,
        "ploshad": rec["ploshad"],
        "tlu": rec["tlu"] or None,
        "god_sozdaniya": str(rec["god"]),
        "metod_sozdaniya": _metod_sozdaniya(rec) or None,
        "sposob_obrabotki": _sposob_obrabotki(rec) or None,
        "glavnaya_poroda": rec["glavnaya_poroda"] or None,
        "sostav_formula": rec["sostav"] or None,
        "gustota_posadki": rec["gustota"],
        "shema_mezhdu_ryadami": rec["shema_mezhdu_ryadami"],
        "shema_v_ryadu": rec["shema_v_ryadu"],
        "posadochnyy_material": rec["material"] or None,
        "istochnik": f"книга л/к: {rec['key']}",
        "primechaniya": " ".join(prim),
    }


def meropriyatiya(rec: dict) -> List[dict]:
    """Записи журнала из строки книги."""
    src = {"istochnik": f"книга л/к: {rec['key']}"}
    note = f"Из книги л/к (лист {rec['list']}, стр. {rec['stroka']})"
    out = []
    if rec["prizhivaemost_1"] is not None:
        out.append(dict(tip=TIP_INV_1, data=str(rec["god"]), prizhivaemost_pct=rec["prizhivaemost_1"],
                        primechaniya=note, dannye=src))
    if rec["prizhivaemost_2"] is not None:
        out.append(dict(tip=TIP_INV_3, data=str(rec["god"] + 2), prizhivaemost_pct=rec["prizhivaemost_2"],
                        primechaniya=note, dannye=src))
    if rec["spisano_ranee"]:
        years = _YEAR_RE.findall(rec["spisano_kniga"])
        out.append(dict(tip=TIP_SPISANIE, data=years[0] if years else "", primechaniya=f"{note}: {rec['spisano_kniga']}",
                        dannye={**src, "prichina": rec["spisano_kniga"]}))
    elif rec["spisat"]:
        years = _YEAR_RE.findall(rec["spisat"])
        card = re.search(r"№\s*(\d+)", rec["spisat"])
        out.append(dict(tip=TIP_SPISANIE, data=years[0] if years else "", primechaniya=f"{note}: {rec['spisat']}",
                        dannye={**src, "prichina": rec["spisat"], "nomer_kartochki": card.group(1) if card else ""}))
    elif rec["dorashchivanie"]:
        years = _YEAR_RE.findall(rec["dorashchivanie"])
        card = re.search(r"№\s*(\d+)", rec["dorashchivanie"])
        out.append(dict(tip=TIP_DORASHCHIVANIE, data=years[0] if years else "",
                        primechaniya=f"{note}: {rec['dorashchivanie']}",
                        dannye={**src, "god_resheniya": years[0] if years else "",
                                "nomer_kartochki": card.group(1) if card else ""}))
    elif rec["perevod_god"]:
        taks = {
            "nomer_kartochki": rec["perevod_kartochka"],
            "ploshad": rec["perevod_ploshad"],
            "sostav": rec["perevod_sostav"],
        }
        out.append(dict(tip=TIP_PEREVOD, data=rec["perevod_god"], sostav_fakt=rec["perevod_sostav"],
                        primechaniya=f"{note}: {rec['perevod_kniga']}",
                        dannye={**src, "taksatsiya": taks}))
    return out


def apply_import(conn, plan: dict, lesnichestvo: str, skip: Iterable[str] = ()) -> dict:
    """Пишет в базу строки плана, кроме ключей из skip. Возвращает счётчики."""
    skip = set(skip)
    created = updated = events = 0
    for rec in plan["rows"]:
        if rec["key"] in skip:
            continue
        fields = uchastok_fields(rec, lesnichestvo)
        if rec["uchastok_id"] is None:
            uchastok_id = legacy_db.create_lesokultury_uchastok(conn, **fields)
            current = {"status": STATUS_ACTIVE}
            created += 1
            have_tips: set = set()
        else:
            uchastok_id = rec["uchastok_id"]
            current = legacy_db.get_lesokultury_uchastok(conn, uchastok_id) or {}
            fill = {}
            for k, v in fields.items():
                if v in (None, ""):
                    continue
                cur = current.get(k)
                if k == "primechaniya":
                    continue  # примечания человека не трогаем
                if cur in (None, ""):
                    fill[k] = v
            # старый скрипт писал в «Выдел» старый номер — меняем на действующий,
            # если его с тех пор не правили руками
            if (rec["vydel_staryy"] and fields["vydel"]
                    and set(vydel_tokens(str(current.get("vydel") or ""))) == set(vydel_tokens(rec["vydel_staryy_kniga"]))
                    and set(vydel_tokens(str(current.get("vydel") or ""))) != set(vydel_tokens(fields["vydel"]))):
                fill["vydel"] = fields["vydel"]
            if fill:
                legacy_db.update_lesokultury_uchastok(conn, uchastok_id, **fill)
                updated += 1
            have_tips = {m["tip"] for m in legacy_db.list_lesokultury_meropriyatiya(conn, uchastok_id)}

        for m in meropriyatiya(rec):
            if m["tip"] in have_tips:
                continue
            legacy_db.add_lesokultury_meropriyatie(
                conn, uchastok_id, m["tip"], m["data"],
                prizhivaemost_pct=m.get("prizhivaemost_pct"),
                sostav_fakt=m.get("sostav_fakt", ""),
                primechaniya=m.get("primechaniya", ""),
                dannye=m.get("dannye"),
            )
            have_tips.add(m["tip"])
            events += 1
        if rec["status"] != STATUS_ACTIVE and (current.get("status") or STATUS_ACTIVE) == STATUS_ACTIVE:
            legacy_db.update_lesokultury_uchastok(conn, uchastok_id, status=rec["status"])
    return {"sozdano": created, "obnovleno": updated, "zapisey_v_zhurnal": events,
            "propushcheno": len([r for r in plan["rows"] if r["key"] in skip])}


def public_rows(rows: List[dict]) -> List[dict]:
    """Строки предпросмотра для сайта (без служебных полей)."""
    keep = ("key", "list", "stroka", "god", "kvartal", "vydel", "vydel_staryy", "vydel_staryy_kniga",
            "vydel_novyy_kniga", "ploshad", "sostav", "status", "deystvie", "uchastok_id", "kak_nayden",
            "perevod_god", "perevod_kartochka", "dorashchivanie", "spisat", "spisano_kniga")
    out = []
    for r in rows:
        d = {k: r.get(k) for k in keep}
        d["problemy"] = r.get("problemy_vse", [])
        out.append(d)
    return out


def parse_overrides(text: Optional[str]) -> Dict[str, str]:
    if not text:
        return {}
    try:
        data = json.loads(text)
    except ValueError:
        raise KnigaError("overrides: ожидается JSON-объект {ключ строки: выдел}")
    if not isinstance(data, dict):
        raise KnigaError("overrides: ожидается JSON-объект {ключ строки: выдел}")
    return {str(k): str(v).strip() for k, v in data.items() if str(v).strip()}
