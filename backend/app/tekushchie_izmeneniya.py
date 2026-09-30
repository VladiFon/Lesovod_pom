# -*- coding: utf-8 -*-
"""
Ведомости текущих изменений (приказ Минлесхоза №130 от 10.06.2026), все
15 приложений:
  - прил. 4, 7, 14 — из «Лесных культур» (перевод в покрытые лесом земли,
    культуры отчётного года, списание);
  - прил. 3, 15 — рубки: выделы делянок (МДО) с актом освидетельствования
    или отметкой «выполнено» в отчётном году, пробы рубок ухода;
  - прил. 5, 6, 8–13 — ручные строки (tek_izm_ruchnye): таких данных в
    программе нет; ручные строки можно добавить и в любое другое приложение;
  - прил. 2 — сводная по всем остальным; прил. 1 — только год и общая
    площадь лесничества (если указана).

Строки собираются из участков и журнала мероприятий; Word заполняется по
шаблону лесничества «Таблицы … ЗАПОЛНЯТЬ ЗДЕСЬ.docx» (лежит в
app/shablony/, можно прислать свой): остальные приложения и шапки шаблона
не трогаются, в них меняются только год и дата заполнения.

Участок культур часто заходит в несколько таксационных выделов и в каждом
становится своим подвыделом — тогда в ведомости по строке на выдел. Такие
части задаются на участке (chasti_json: [{vydel, podvydel, ploshad}]);
если их нет и выдел один, строка берётся из самого участка.
"""
import copy
import datetime as dt
import io
import json
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from app import legacy_bridge  # noqa: F401
from app import polnota as polnota_mod
from app.lesokultury_kniga import vydel_tokens

TEMPLATE_PATH = Path(__file__).resolve().parent / "shablony" / "tekushchie_izmeneniya.docx"

TIP_PEREVOD = "Перевод в покрытые лесом земли"
TIP_SPISANIE = "Списание"
TIPY_PRIZHIVAEMOSTI = ("Инвентаризация 1-го года", "Техническая приёмка")
DIAMETR_KULTUR_SM = 2  # у переводимых несомкнувшихся культур (со слов лесничества)

# Заголовки столбцов — как в формах приказа (для предпросмотра и ручного ввода).
_TAKS = ["Кв.", "Выдел (подвыдел) по лесоустройству", "Площадь, га", "Выдел (подвыдел)", "Площадь, га",
         "Состав", "Возраст, лет", "Высота, м", "Диаметр, см", "Полнота"]
_KULT = ["Кв.", "Выдел по лесоустройству", "Площадь, га", "Выдел (подвыдел)", "Площадь, га",
         "Метод, способ создания", "Состав", "Обработка почвы", "Между рядами, м", "В ряду, м",
         "Количество, шт/га", "Приживаемость, %"]
COLUMNS = {
    3: ["Кв.", "Выдел по лесоустройству", "Площадь, га", "Вид рубки", "Выдел (подвыдел)",
        "Площадь вырубки, га", "Выбираемый запас, м3/га"],
    4: _TAKS,
    5: _TAKS,
    6: _TAKS,
    7: _KULT,
    8: _KULT[:5] + ["Назначение плантации"] + _KULT[6:],
    9: ["Кв.", "Выдел по лесоустройству", "Площадь, га", "Метод естественного возобновления",
        "Выдел (подвыдел)", "Площадь, га", "Метод содействия"],
    10: _TAKS,
    11: ["Кадастровый номер", "Кв.", "Выдел", "Категория леса", "Район", "Ограничение режима",
         "Площадь, га", "Вид земель", "Состав", "Возраст, лет", "Высота, м", "Диаметр, см", "Полнота",
         "Тип леса"],
    12: ["Кадастровый номер", "Кв.", "Выдел (подвыдел)", "Площадь, га"],
    13: ["Кв.", "Выдел по лесоустройству", "Площадь, га", "Выдел (подвыдел)", "Площадь, га",
         "Новый вид земель", "Примечание"],
    14: ["Кв.", "Выдел", "Площадь, га", "Вид земель", "Состав", "Возраст, лет", "Высота, м",
         "Диаметр, см", "Полнота", "Тип леса", "Причина перевода"],
    15: ["Кв.", "Выдел по лесоустройству", "Площадь, га", "Вид рубки", "Выдел (подвыдел)", "Площадь, га",
         "Полнота после рубки", "Выбираемый запас, м3/га"],
}
TITLES = {
    3: "Прил. 3 — сплошнолесосечные, сплошные санитарные рубки, рубки реконструкции, окончательные приёмы постепенных",
    4: "Прил. 4 — несомкнувшиеся лесные культуры, переведённые в покрытые лесом земли",
    5: "Прил. 5 — участки с содействием естественному возобновлению, переведённые в покрытые лесом земли",
    6: "Прил. 6 — участки естественного возобновления, переведённые в покрытые лесом земли",
    7: "Прил. 7 — лесные культуры, созданные в отчётном году",
    8: "Прил. 8 — плантации и объекты постоянной лесосеменной базы",
    9: "Прил. 9 — содействие естественному возобновлению",
    10: "Прил. 10 — мягколиственные насаждения, введённые в категорию ценных рубками ухода",
    11: "Прил. 11 — земли, предоставленные в состав лесного фонда",
    12: "Прил. 12 — земли, изъятые из состава лесного фонда",
    13: "Прил. 13 — лесные земли, переведённые в нелесные",
    14: "Прил. 14 — участки, переведённые из одного вида земель в другой (списание культур)",
    15: "Прил. 15 — несплошные рубки главного пользования, рубки промежуточного пользования, прочие рубки",
}
# Откуда берутся строки (подсказка на сайте). Ручные строки можно добавить в любое приложение.
ISTOCHNIKI = {
    3: "делянки (МДО) с актом освидетельствования или отметкой «выполнено» в отчётном году",
    4: ("лесные культуры: «Перевод в покрытые лесом земли» в журнале участка; таксацию можно загрузить "
        "из полевых карточек («Лесокультуры» → «Загрузить карточки перевода»)"),
    7: "лесные культуры, созданные в отчётном году",
    14: "лесные культуры: «Списание» в журнале участка",
    15: "делянки (МДО) с несплошной рубкой и выполненные пробы рубок ухода",
}
NOMERA = tuple(sorted(COLUMNS))
# Столбец площади (для сводной прил. 2): изменённая характеристика, если она есть.
PLOSHAD_COL = {3: 5, 4: 4, 5: 4, 6: 4, 7: 4, 8: 4, 9: 5, 10: 4, 11: 6, 12: 3, 13: 4, 14: 2, 15: 5}

# Виды рубок (для прил. 3/15 и строк прил. 2).
RUBKA_SPLOSHNAYA_GLAVNAYA = "сплошнолесосечная рубка главного пользования"
RUBKA_SPLOSHNAYA_SANITARNAYA = "сплошная санитарная рубка"
RUBKA_SPLOSHNAYA_REKONSTRUKTSII = "сплошная рубка реконструкции"
RUBKA_OKONCHATELNYY_PRIEM = "окончательный приём постепенной рубки"
RUBKA_NESPLOSHNAYA_GLAVNAYA = "несплошная рубка главного пользования"
RUBKA_PROMEZHUTOCHNAYA = "рубка промежуточного пользования"
RUBKA_PROCHAYA = "прочая рубка"

# Строки прил. 2: (номер строки таблицы, откуда считать).
SVODNAYA = {
    2: ("rubka", 3, RUBKA_SPLOSHNAYA_GLAVNAYA),
    3: ("rubka", 3, RUBKA_SPLOSHNAYA_SANITARNAYA),
    4: ("rubka", 3, RUBKA_SPLOSHNAYA_REKONSTRUKTSII),
    5: ("rubka", 3, RUBKA_OKONCHATELNYY_PRIEM),
    6: ("sum", (7, 8, 9)),
    7: ("pril", 4),
    8: ("pril", 5),
    9: ("pril", 6),
    10: ("sum", (11, 12, 13)),
    11: ("kultury", "posadka"),
    12: ("kultury", "posev"),
    13: ("kultury", "rekonstruktsiya"),
    14: ("pril", 8),
    15: ("pril", 9),
    16: ("pril", 10),
    17: ("pril", 11),
    18: ("pril", 12),
    19: ("pril", 13),
    20: ("pril", 14),
    22: ("rubka", 15, RUBKA_NESPLOSHNAYA_GLAVNAYA),
    23: ("rubka", 15, RUBKA_PROMEZHUTOCHNAYA),
    24: ("rubka", 15, RUBKA_PROCHAYA),
}


class TIError(Exception):
    """Ошибка, понятная человеку (показывается на сайте как есть)."""


# --------------------------------------------------------------------------- #
#   Сбор строк
# --------------------------------------------------------------------------- #
def _norm(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower().replace("ё", "е")


def _year_of(text) -> Optional[int]:
    found = re.findall(r"(?:19|20)\d{2}", str(text or ""))
    return int(found[-1]) if found else None


def fmt(value) -> str:
    """Число без лишних нулей: 4421.0 -> «4421», 0.60 -> «0.6»; текст как есть."""
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        text = f"{float(value):.3f}".rstrip("0").rstrip(".")
        return text if text != "-0" else "0"
    return str(value).strip()


def _cap(text) -> str:
    text = str(text or "").strip()
    return text[:1].upper() + text[1:] if text else ""


def _num(value) -> Optional[float]:
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None


def _taxation_ploshad(conn, lesnichestvo: str) -> Dict[Tuple[str, str], str]:
    """{(квартал, выдел): площадь по лесоустройству} для лесничества."""
    target = _norm(lesnichestvo)
    try:
        rows = conn.execute(
            "SELECT l.name, k.nomer, v.nomer, v.ploshad FROM vydel v "
            "JOIN kvartal k ON k.id = v.kvartal_id JOIN lesnichestvo l ON l.id = k.lesnichestvo_id"
        ).fetchall()
    except Exception:  # noqa: BLE001 — таблиц таксации может не быть
        return {}
    result = {}
    for name, kv, vd, pl in rows:
        if not _lesn_match(target, name):
            continue
        result[(str(kv).strip(), str(vd).strip())] = fmt(_num(pl)) if _num(pl) is not None else str(pl or "")
    return result


def _uchastki(conn, lesnichestvo: str) -> List[dict]:
    conn_rows = conn.execute("SELECT * FROM lesokultury_uchastok ORDER BY id").fetchall()
    cols = [d[0] for d in conn.execute("SELECT * FROM lesokultury_uchastok LIMIT 0").description]
    target = _norm(lesnichestvo)
    out = []
    for row in conn_rows:
        u = dict(zip(cols, row))
        if not _lesn_match(target, u.get("lesnichestvo")):
            continue
        out.append(u)
    return out


def _meropriyatiya(conn) -> Dict[int, List[dict]]:
    rows = conn.execute(
        "SELECT uchastok_id, tip, data, prizhivaemost_pct, sostav_fakt, primechaniya, dannye_json, id "
        "FROM lesokultury_meropriyatiya ORDER BY id"
    ).fetchall()
    result: Dict[int, List[dict]] = {}
    for uid, tip, data, pct, sostav, prim, dannye, mid in rows:
        try:
            parsed = json.loads(dannye) if dannye else {}
        except ValueError:
            parsed = {}
        result.setdefault(uid, []).append({
            "id": mid, "tip": tip, "data": data, "prizhivaemost_pct": pct,
            "sostav_fakt": sostav or "", "primechaniya": prim or "", "dannye": parsed or {},
        })
    return result


def chasti(u: dict) -> List[dict]:
    """Части участка по выделам: [{vydel, podvydel, ploshad}]. Без заданных
    частей — одна часть из самого участка (для нескольких выделов — по
    части на выдел, площадь не известна)."""
    raw = u.get("chasti_json")
    if raw:
        try:
            parts = json.loads(raw)
        except ValueError:
            parts = []
        clean = [
            {"vydel": str(p.get("vydel") or "").strip(), "podvydel": str(p.get("podvydel") or "").strip(),
             "ploshad": _num(p.get("ploshad"))}
            for p in parts if isinstance(p, dict) and (p.get("vydel") or p.get("podvydel"))
        ]
        if clean:
            return clean
    tokens = vydel_tokens(str(u.get("vydel") or ""))
    if len(tokens) <= 1:
        return [{"vydel": tokens[0] if tokens else "", "podvydel": str(u.get("podvydel") or "").strip(),
                 "ploshad": _num(u.get("ploshad"))}]
    return [{"vydel": t, "podvydel": "", "ploshad": None} for t in tokens]


def _mesto(u: dict) -> str:
    return f"кв. {u.get('kvartal') or '—'} выд. {u.get('vydel') or '—'}"


def _chasti_warning(u: dict, parts: List[dict]) -> Optional[str]:
    if len(parts) > 1 and any(p["ploshad"] is None for p in parts):
        return (f"{_mesto(u)}: участок в нескольких выделах — нажмите «По выделам» и укажите "
                "подвыдел и площадь в каждом")
    return None


def resolve_lesnichestvo(value) -> str:
    """Номер лесничества (как его хранят настройки входа: «5») -> название
    из LCH_MAP («Болбасовское»); название возвращается как есть."""
    text = str(value or "").strip()
    if not text.isdigit():
        return text
    try:
        import config as legacy_config
        for name, num in legacy_config.LCH_MAP.items():
            if str(num).strip() == text:
                return name
    except Exception:  # noqa: BLE001
        pass
    return text


def _lesn_match(target: str, own) -> bool:
    """Лесничество совпадает, если одно название содержит другое или у них
    общая основа первого слова («Болбасовское», «Болбасовского л-ва»)."""
    own = _norm(own)
    if not target or not own or own == target or target in own or own in target:
        return True
    a, b = target.split()[0], own.split()[0]
    n = min(len(a), len(b), 7)
    return n >= 5 and a[:n] == b[:n]


def _row(res: dict, key: str, row: list, vid: Optional[str] = None, dop: str = "") -> None:
    """Автоматическая строка с ключом (по нему хранятся дописанные вручную
    значения, см. popravki) и, для прил. 3/7/15, видом для сводной. dop —
    подсказка только для экрана (в Word не идёт): год культур, № карточки."""
    res["rows"].append(row)
    res["keys"].append(key)
    res.setdefault("dop", []).append(dop)
    if vid is not None:
        res.setdefault("vidy", []).append(vid)


def _chasti_info(res: dict, prefix: str, kv: str, vydel, ploshad, raw, taxation: dict) -> None:
    """Для кнопки «По выделам» / «Новый выдел»: исходный выдел(ы), общая
    площадь, заданные части и площади выделов по лесоустройству. Номер,
    который можно дать новому выделу, дописывает _novye_vydely()."""
    tokens = vydel_tokens(str(vydel or ""))
    parsed = []
    if raw:
        try:
            parsed = [p for p in json.loads(raw) if isinstance(p, dict)]
        except (TypeError, ValueError):
            parsed = []
    res.setdefault("chasti_info", {})[prefix] = {
        "kvartal": kv, "vydel": str(vydel or ""), "ploshad": fmt(_num(ploshad)),
        "vydely": tokens, "chasti": parsed,
        "taks": {t: taxation.get((kv, t), "") for t in tokens},
    }


def _warn(res: dict, key: Optional[str], text: str, cols: Optional[List[int]] = None) -> None:
    """Замечание; key — ключ строки или её начало (u12 — все строки участка
    12), cols — графы, о которых оно. Когда эти графы (по умолчанию —
    NUZHNYE) во всех таких строках дописаны, замечание не показывается."""
    res["warn_keys"].append((key, text, cols))


# Поля участка, по которым видно, насколько он заполнен (для выбора
# основного из дублей).
_POLYA_POLNOTY = ("metod_sozdaniya", "sostav_formula", "sposob_obrabotki", "shema_mezhdu_ryadami",
                  "shema_v_ryadu", "gustota_posadki", "podvydel", "chasti_json", "geom_geojson", "tlu",
                  "glavnaya_poroda", "posadochnyy_material")


def _dubl_klyuch(u: dict):
    """Одно и то же место: квартал, набор выделов и год создания. Без
    выдела или года участки дублями не считаются."""
    tokens = set(vydel_tokens(str(u.get("vydel") or "")))
    tokens |= {p["vydel"] for p in chasti(u) if p["vydel"]}
    god_s = _year_of(u.get("god_sozdaniya"))
    kv = _norm(u.get("kvartal"))
    if not tokens or not god_s or not kv:
        return None
    return kv, frozenset(tokens), god_s


def _zapolnennost(conn, u: dict, journal: List[dict]) -> int:
    score = sum(1 for f in _POLYA_POLNOTY if str(u.get(f) or "").strip())
    score += 2 * len(journal)
    ensure_popravki(conn)
    score += 3 * conn.execute("SELECT COUNT(*) FROM tek_izm_popravki WHERE klyuch LIKE ?",
                              (f"u{u['id']}:%",)).fetchone()[0]
    return score


def dubli(conn, uchastki: List[dict], mer: Dict[int, List[dict]]) -> Tuple[List[dict], List[dict]]:
    """Участки, заведённые дважды (например, из книги л/к и ещё раз вручную
    или из QGIS): в ведомость идёт только самый заполненный, остальные —
    в список дублей (на сайте кнопка «Объединить»). Возвращает
    (участки без дублей, [{osnovnoy, lishnie, mesto}])."""
    groups: Dict[tuple, List[dict]] = {}
    for u in uchastki:
        key = _dubl_klyuch(u)
        if key is None:
            continue
        lesn = _norm(u.get("lesnichestvo"))
        for (k, other_lesn), members in groups.items():
            if k == key and (not lesn or not other_lesn or _lesn_match(lesn, other_lesn)):
                members.append(u)
                break
        else:
            groups[(key, lesn)] = [u]
    skip, out = set(), []
    rank = lambda group: sorted(group, key=lambda u: (-_zapolnennost(conn, u, mer.get(u["id"], [])), u["id"]))
    for (key, _lesn), members in groups.items():
        # Разные строки книги л/к в одном выделе — разные участки (посадки
        # разных лет/частей), это не дубль. Дубль — участок, заведённый
        # вручную или из QGIS поверх уже существующего.
        iz_knigi = [u for u in members if str(u.get("istochnik") or "").startswith("книга")]
        drugie = [u for u in members if u not in iz_knigi]
        if len(members) < 2 or not drugie:
            continue
        if len(iz_knigi) <= 1:
            ranked = rank(members)
        else:
            ranked = rank(iz_knigi)[:1] + rank(drugie)
        main = ranked[0]
        skip.update(u["id"] for u in ranked[1:])
        out.append({
            "osnovnoy": main["id"],
            "lishnie": [u["id"] for u in ranked[1:]],
            "mesto": f"{_mesto(main)} ({key[2]} г.)",
            "uchastki": [{"id": u["id"], "vydel": u.get("vydel") or "", "ploshad": fmt(_num(u.get("ploshad"))),
                          "primechaniya": u.get("primechaniya") or "", "created_at": u.get("created_at") or ""}
                         for u in ranked],
        })
    return [u for u in uchastki if u["id"] not in skip], out


def _kultury(conn, god: int, lesnichestvo: str, result: dict, taxation: dict) -> None:
    """Прил. 4, 7, 14 из «Лесных культур»."""
    mer = _meropriyatiya(conn)
    uchastki, result["dubli"] = dubli(conn, _uchastki(conn, lesnichestvo), mer)
    for u in uchastki:
        kv = str(u.get("kvartal") or "").strip()
        parts = chasti(u)
        journal = mer.get(u["id"], [])

        # --- Прил. 7: создан в отчётном году
        if _year_of(u.get("god_sozdaniya")) == god:
            res = result[7]
            res["uchastki"] += 1
            _chasti_info(res, f"u{u['id']}", kv, u.get("vydel"), u.get("ploshad"), u.get("chasti_json"), taxation)
            pct = next((m["prizhivaemost_pct"] for m in reversed(journal)
                        if m["tip"] in TIPY_PRIZHIVAEMOSTI and m["prizhivaemost_pct"] is not None), None)
            missing = [name for name, val in (
                ("способ обработки почвы", u.get("sposob_obrabotki")),
                ("схема посадки", u.get("shema_mezhdu_ryadami") and u.get("shema_v_ryadu")),
                ("подвыдел", all(p["podvydel"] for p in parts)),
                ("приживаемость (инвентаризация 1-го года)", pct is not None),
            ) if not val]
            if missing:
                _warn(res, f"u{u['id']}", f"{_mesto(u)}: не заполнено — " + ", ".join(missing), [3, 5, 7, 8, 9, 11])
            warn = _chasti_warning(u, parts)
            if warn:
                _warn(res, f"u{u['id']}", warn, [4])
            for i, p in enumerate(parts):
                _row(res, f"u{u['id']}:{i}", [
                    kv, p["vydel"], taxation.get((kv, p["vydel"]), ""),
                    p["podvydel"] or p["vydel"], fmt(p["ploshad"]),
                    _cap(u.get("metod_sozdaniya")), u.get("sostav_formula") or "",
                    _cap(u.get("sposob_obrabotki")), fmt(u.get("shema_mezhdu_ryadami")),
                    fmt(u.get("shema_v_ryadu")), fmt(u.get("gustota_posadki")), fmt(pct),
                ], _vid_kultur(u.get("metod_sozdaniya")))

        # --- Прил. 4: перевод в отчётном году
        perevod = [m for m in journal if m["tip"] == TIP_PEREVOD and _year_of(m["data"]) == god]
        if perevod:
            m = perevod[-1]
            t = m["dannye"].get("taksatsiya") or {}
            res = result[4]
            res["uchastki"] += 1
            _chasti_info(res, f"u{u['id']}", kv, u.get("vydel"), u.get("ploshad"), u.get("chasti_json"), taxation)
            # Возраст культур — от года создания, если при переводе не записан.
            vozrast = t.get("vozrast")
            if vozrast in (None, "") and _year_of(u.get("god_sozdaniya")):
                vozrast = god - _year_of(u.get("god_sozdaniya"))
            sostav = t.get("sostav") or m["sostav_fakt"] or u.get("sostav_formula") or ""
            # Диаметр у переводимых культур в карточке не пишется — всегда 2 см.
            diametr = t.get("diametr") if t.get("diametr") not in (None, "") else DIAMETR_KULTUR_SM
            # Полнота — по таблице молодняков из количества деревьев на 1 га.
            polnota = t.get("polnota")
            if polnota in (None, "") and _num(t.get("kolichestvo_tys_na_ga")):
                polnota, _nizhe = polnota_mod.po_kolichestvu(_num(t.get("kolichestvo_tys_na_ga")), sostav)
            if t.get("polnota_nizhe_tablicy"):
                _warn(res, None, f"{_mesto(u)}: деревьев {fmt(t.get('kolichestvo_tys_na_ga'))} тыс./га — меньше, чем в "
                                 f"таблице полноты; поставлено {fmt(polnota)}, проверьте")
            missing = [name for name, val in (("возраст", vozrast), ("высота", t.get("vysota")),
                                              ("полнота", polnota))
                       if val in (None, "")]
            if missing:
                _warn(res, f"u{u['id']}", f"{_mesto(u)}: в таксации при переводе нет — " + ", ".join(missing),
                      [6, 7, 8, 9])
            # Выдел(ы) перевода — по полевой карточке, если она загружена и
            # части по выделам на участке не заданы вручную.
            vk = vydel_tokens(str(t.get("vydel_kartochki") or ""))
            if vk and not u.get("chasti_json"):
                ploshad_k = _num(t.get("ploshad")) or _num(u.get("ploshad"))
                parts = ([{"vydel": vk[0], "podvydel": "", "ploshad": ploshad_k}] if len(vk) == 1
                         else [{"vydel": x, "podvydel": "", "ploshad": None} for x in vk])
            # Несколько старых выделов стали одним новым («6,7,8,9 (30)» в
            # карточке) и площади по выделам не заданы — одна строка.
            if len(parts) > 1 and t.get("podvydel") and all(p["ploshad"] is None for p in parts):
                parts = [{"vydel": ", ".join(p["vydel"] for p in parts), "podvydel": "",
                          "ploshad": _num(t.get("ploshad")) or _num(u.get("ploshad"))}]
            warn = _chasti_warning(u, parts)
            if warn:
                _warn(res, f"u{u['id']}", warn, [4])
            one = len(parts) == 1
            god_k = _year_of(u.get("god_sozdaniya"))
            dop = ", ".join(x for x in (
                f"{god_k} г." if god_k else "",
                f"карт. №{t['nomer_kartochki']}" if t.get("nomer_kartochki") else "",
            ) if x)
            for i, p in enumerate(parts):
                # Новый выдел: заданный вручную («По выделам» / «Новый
                # выдел»), иначе из карточки (в скобках), иначе прежний.
                novyy = p["podvydel"] or (t.get("podvydel") if one and t.get("podvydel") else "") or p["vydel"]
                _row(res, f"u{u['id']}:{i}", [
                    kv, p["vydel"], fmt(p["ploshad"]), novyy,
                    fmt(_num(t.get("ploshad")) if one and t.get("ploshad") and not u.get("chasti_json")
                        else p["ploshad"]),
                    sostav, fmt(_num(vozrast) if _num(vozrast) is not None else vozrast), fmt(t.get("vysota")),
                    fmt(diametr), fmt(polnota),
                ], dop=dop)

        # --- Прил. 14: списание в отчётном году
        spisanie = [m for m in journal if m["tip"] == TIP_SPISANIE and _year_of(m["data"]) == god]
        if spisanie:
            m = spisanie[-1]
            d = m["dannye"]
            res = result[14]
            res["uchastki"] += 1
            _chasti_info(res, f"u{u['id']}", kv, u.get("vydel"), u.get("ploshad"), u.get("chasti_json"), taxation)
            if not d.get("vid_zemel"):
                _warn(res, f"u{u['id']}", f"{_mesto(u)}: не указан вид земель после списания", [3])
            warn = _chasti_warning(u, parts)
            if warn:
                _warn(res, f"u{u['id']}", warn, [4])
            god_s = _year_of(u.get("god_sozdaniya"))
            prichina = d.get("prichina") or m["primechaniya"] or ""
            akt = " ".join(x for x in (
                f"акт №{d['akt_nomer']}" if d.get("akt_nomer") else "",
                f"от {d['akt_data']}" if d.get("akt_data") else "",
            ) if x)
            for i, p in enumerate(parts):
                _row(res, f"u{u['id']}:{i}", [
                    kv, p["podvydel"] or p["vydel"], fmt(p["ploshad"]), d.get("vid_zemel") or "",
                    u.get("sostav_formula") or "", fmt(god - god_s) if god_s else "", "", "", "", "",
                    "; ".join(x for x in (prichina, akt) if x),
                ])



def _vid_kultur(metod) -> str:
    """Для прил. 2: посадка / посев / реконструкция (по методу создания)."""
    m = _norm(metod)
    if "реконстр" in m:
        return "rekonstruktsiya"
    if "посев" in m or "аэросев" in m:
        return "posev"
    return "posadka"


# --------------------------------------------------------------------------- #
#   Рубки: прил. 3 и 15
# --------------------------------------------------------------------------- #
_PROMEZH = ("промежут", "уход", "осветл", "прочист", "прорежив", "проходн", "санитар", "ландшафт",
            "обновлен", "переформир", "реконстр")
_NESPL_GLAV = ("главн", "постепен", "выборочн", "группов", "длительно")


def vid_rubki(*texts) -> Optional[str]:
    """Вид рубки по тексту МДО/ведомости («рубка леса», «вид рубки») — одна
    из констант RUBKA_*; None, если текста нет."""
    t = _norm(" ".join(str(x or "") for x in texts))
    if not t:
        return None
    if "сплош" in t and "несплош" not in t:
        if "санитар" in t:
            return RUBKA_SPLOSHNAYA_SANITARNAYA
        if "реконстр" in t:
            return RUBKA_SPLOSHNAYA_REKONSTRUKTSII
        return RUBKA_SPLOSHNAYA_GLAVNAYA
    if "постепен" in t and "оконч" in t:
        return RUBKA_OKONCHATELNYY_PRIEM
    if "главн" in t:
        return RUBKA_NESPLOSHNAYA_GLAVNAYA
    if "проч" in t:
        return RUBKA_PROCHAYA
    if any(k in t for k in _PROMEZH):
        return RUBKA_PROMEZHUTOCHNAYA
    if any(k in t for k in _NESPL_GLAV):
        return RUBKA_NESPLOSHNAYA_GLAVNAYA
    return None


def pril_rubki(vid: Optional[str]) -> int:
    return 3 if vid in (RUBKA_SPLOSHNAYA_GLAVNAYA, RUBKA_SPLOSHNAYA_SANITARNAYA,
                        RUBKA_SPLOSHNAYA_REKONSTRUKTSII, RUBKA_OKONCHATELNYY_PRIEM) else 15


def _json(text) -> dict:
    try:
        value = json.loads(text) if text else {}
    except (TypeError, ValueError):
        return {}
    return value if isinstance(value, dict) else {}


def _rubki_delyanki(conn, god: int, lesnichestvo: str, result: dict, taxation: dict) -> None:
    """Выделы делянок, вырубленные в отчётном году: акт освидетельствования
    делянки за этот год или выдел отмечен «выполнено» с датой в этом году."""
    target = _norm(lesnichestvo)
    try:
        items = conn.execute(
            "SELECT i.id, i.delyanka_id, d.nazvanie, i.lesnichestvo, i.kvartal, i.vydel, i.ploshad, "
            "i.zapas_na_ga, i.vyrubaemyy_zapas, i.polnota, i.mdo_raw_json, i.status_rabot, i.data_vypolneniya, "
            "i.chasti_json "
            "FROM delyanka_item i JOIN delyanka d ON d.id = i.delyanka_id ORDER BY i.id"
        ).fetchall()
        acts = conn.execute("SELECT delyanka_id, act_date FROM osvidetelstvovanie_acts").fetchall()
    except Exception:  # noqa: BLE001 — старая база без делянок
        return
    act_years: Dict[int, set] = {}
    for d_id, act_date in acts:
        y = _year_of(act_date)
        if y:
            act_years.setdefault(d_id, set()).add(y)

    for (item_id, d_id, nazvanie, lesn, kv, vd, pl, zapas_ga, vyrub, polnota, raw, status,
         data_vyp, chasti_raw) in items:
        if not _lesn_match(target, lesn):
            continue
        done = _norm(status) == "выполнено"
        y_vyp = _year_of(data_vyp) if done else None
        if not (god in act_years.get(d_id, set()) or y_vyp == god):
            if done and y_vyp is None and not act_years.get(d_id):
                mesto = f"кв. {kv or '—'} выд. {vd or '—'}"
                _warn(result[3], None,
                    f"{mesto} (делянка «{nazvanie or d_id}»): отмечен выполненным, но нет ни даты, ни акта "
                    "освидетельствования — в ведомости рубок не учтён")
            continue
        mdo = _json(raw)
        vid = vid_rubki(mdo.get("vid_rubki"), mdo.get("sposob_rubki"))
        n = pril_rubki(vid)
        res = result[n]
        res["uchastki"] += 1
        kv = str(kv or "").strip()
        vd = str(vd or "").strip()
        mesto = f"кв. {kv or '—'} выд. {vd or '—'}"
        if vid is None:
            _warn(res, f"d{item_id}", f"{mesto}: в МДО не указан вид рубки — отнесён к прочим рубкам", [3])
            vid = RUBKA_PROCHAYA
        area = _num(pl)
        zapas = _num(vyrub)
        if zapas is not None and area:
            zapas_ga = zapas / area
        else:
            zapas_ga = _num(zapas_ga)
            vyborka = _num(str(mdo.get("vyborka_zapasa_pct") or "").replace("%", ""))
            if zapas_ga is not None and vyborka and n == 15:
                zapas_ga = zapas_ga * vyborka / 100
        if zapas_ga is None:
            _warn(res, f"d{item_id}", f"{mesto}: нет выбираемого запаса (МДО)", [6 if n == 3 else 7])
        nazvanie_vida = _cap(mdo.get("sposob_rubki") or mdo.get("vid_rubki") or vid)
        posle = None
        if n == 15:
            p0 = _num(polnota) or _num(mdo.get("polnota"))
            vyborka = _num(str(mdo.get("vyborka_zapasa_pct") or "").replace("%", ""))
            posle = round(p0 * (1 - vyborka / 100), 1) if p0 and vyborka else None
            if posle is None:
                _warn(res, f"d{item_id}", f"{mesto}: полноту после рубки не из чего посчитать (нет полноты или % выборки)",
                      [6])
        parts = chasti({"vydel": vd, "ploshad": area, "chasti_json": chasti_raw})
        _chasti_info(res, f"d{item_id}", kv, vd, area, chasti_raw, taxation)
        if len(parts) > 1 and any(p["ploshad"] is None for p in parts):
            _warn(res, f"d{item_id}", f"{mesto}: лесосека в нескольких выделах — нажмите «По выделам» и "
                                      "укажите площадь в каждом", [5])
        for i, p in enumerate(parts):
            pv = p["vydel"] or vd
            tax = taxation.get((kv, pv), "") or (fmt(_num(mdo.get("ploshad_obshaya"))) if len(parts) == 1 else "")
            row = [kv, pv, tax, nazvanie_vida, p["podvydel"] or pv, fmt(p["ploshad"])]
            if n == 15:
                row.append(fmt(posle))
            row.append(fmt(round(zapas_ga)) if zapas_ga else "")
            _row(res, f"d{item_id}:{i}", row, vid)


def _rubki_uhoda(conn, god: int, lesnichestvo: str, result: dict, taxation: dict) -> None:
    """Пробы рубок ухода, отмеченные выполненными в отчётном году → прил. 15."""
    target = _norm(lesnichestvo)
    try:
        rows = conn.execute(
            "SELECT id, kvartal, vydel, ploshad_vydela, data_zamera, data_json, completed_at FROM uhody_proby "
            "WHERE completed_at IS NOT NULL AND completed_at != '' ORDER BY id"
        ).fetchall()
    except Exception:  # noqa: BLE001
        return
    res = result[15]
    for pid, kv, vd, pl_vydela, data_zamera, data_json, completed in rows:
        if _year_of(completed) != god:
            continue
        d = _json(data_json)
        if not _lesn_match(target, d.get("lesnichestvo")):
            continue
        kv = str(kv or "").strip()
        vd = str(vd or "").strip()
        mesto = f"кв. {kv or '—'} выд. {vd or '—'}"
        vid = vid_rubki(d.get("vid_polzovaniya"), d.get("vid_rubki"), d.get("sposob_rubki")) or RUBKA_PROMEZHUTOCHNAYA
        if pril_rubki(vid) == 3:
            vid = RUBKA_PROMEZHUTOCHNAYA
        res["uchastki"] += 1
        area = _num(d.get("ploshad_lesoseki")) or _num(pl_vydela)
        zapas_ga = _num(d.get("zapas_na_1ga"))
        polnota = _num(d.get("polnota"))
        if polnota is None:
            _warn(res, f"p{pid}", f"{mesto} (проба рубок ухода №{pid}): не указана полнота", [6])
        tax = taxation.get((kv, vd), "") or fmt(_num(pl_vydela))
        _row(res, f"p{pid}", [
            kv, vd, tax, _cap(d.get("vid_rubki") or d.get("sposob_rubki") or "рубка ухода"), vd,
            fmt(area), fmt(polnota), fmt(round(zapas_ga, 1)) if zapas_ga else "",
        ], vid)


# --------------------------------------------------------------------------- #
#   Ручные строки (прил. 5, 6, 8–13 — данных для них в программе нет; и
#   дополнения к любому другому приложению)
# --------------------------------------------------------------------------- #
def ensure_table(conn) -> None:
    conn.execute(
        "CREATE TABLE IF NOT EXISTS tek_izm_ruchnye ("
        " id INTEGER PRIMARY KEY AUTOINCREMENT,"
        " god INTEGER NOT NULL,"
        " lesnichestvo TEXT DEFAULT '',"
        " prilozhenie INTEGER NOT NULL,"
        " znacheniya_json TEXT NOT NULL,"
        " created_at TEXT DEFAULT (datetime('now', 'localtime')))"
    )


def ruchnye(conn, god: int, lesnichestvo: str = "") -> List[dict]:
    ensure_table(conn)
    target = _norm(lesnichestvo)
    out = []
    for rid, pril, lesn, values in conn.execute(
        "SELECT id, prilozhenie, lesnichestvo, znacheniya_json FROM tek_izm_ruchnye WHERE god = ? ORDER BY id",
        (god,),
    ).fetchall():
        if not _lesn_match(target, lesn):
            continue
        try:
            parsed = json.loads(values)
        except ValueError:
            parsed = []
        cols = len(COLUMNS.get(pril, []))
        parsed = [str(v if v is not None else "").strip() for v in (parsed if isinstance(parsed, list) else [])]
        out.append({"id": rid, "prilozhenie": pril, "lesnichestvo": lesn or "",
                    "values": (parsed + [""] * cols)[:cols]})
    return out


def _ruchnye_v_result(conn, god: int, lesnichestvo: str, result: dict) -> None:
    for r in ruchnye(conn, god, lesnichestvo):
        n = r["prilozhenie"]
        if n not in result:
            continue
        res = result[n]
        res["rows"].append(r["values"])
        res["keys"].append(f"r{r['id']}")
        res["pustye"].append([])
        res["popravleno"].append(False)
        res.setdefault("dop", []).append("")
        res["ruchnye"].append({"id": r["id"], "values": r["values"]})
        if n in (3, 15):
            vid = vid_rubki(r["values"][3])
            if vid is None or pril_rubki(vid) != n:
                vid = RUBKA_SPLOSHNAYA_GLAVNAYA if n == 3 else RUBKA_PROCHAYA
            res.setdefault("vidy", []).append(vid)
        elif n == 7:
            res.setdefault("vidy", []).append(_vid_kultur(r["values"][5]))


# --------------------------------------------------------------------------- #
#   Сводная (прил. 2)
# --------------------------------------------------------------------------- #
def _itog(rows: List[list], col: int) -> Tuple[float, int]:
    area = 0.0
    for r in rows:
        v = _num(r[col]) if col < len(r) else None
        area += v or 0.0
    return round(area, 2), len(rows)


def svodnaya(result: dict) -> Dict[int, Tuple[float, int]]:
    """{номер строки таблицы прил. 2: (площадь, количество участков)}."""
    out: Dict[int, Tuple[float, int]] = {}
    for row_no, rule in SVODNAYA.items():
        kind = rule[0]
        if kind == "pril":
            n = rule[1]
            out[row_no] = _itog(result[n]["rows"], PLOSHAD_COL[n])
        elif kind in ("rubka", "kultury"):
            n = rule[1] if kind == "rubka" else 7
            want = rule[2] if kind == "rubka" else rule[1]
            vidy = result[n].get("vidy", [])
            rows = [r for r, v in zip(result[n]["rows"], vidy) if v == want]
            out[row_no] = _itog(rows, PLOSHAD_COL[n])
    for row_no, rule in SVODNAYA.items():
        if rule[0] == "sum":
            parts = [out[i] for i in rule[1]]
            out[row_no] = (round(sum(p[0] for p in parts), 2), sum(p[1] for p in parts))
    return out


def diagnostika(conn, god: int, lesnichestvo: str = "") -> dict:
    """Что вообще есть в базе — показывается на сайте, когда ведомости
    пустые: какие лесничества и годы записаны у культур и делянок, сколько
    переводов/списаний и актов по годам. По ней видно, что не совпало
    (лесничество, год, статус)."""
    target = _norm(lesnichestvo)

    def safe(sql, params=()):
        try:
            return conn.execute(sql, params).fetchall()
        except Exception:  # noqa: BLE001
            return []

    def counter(values) -> List[list]:
        out: Dict[str, int] = {}
        for v in values:
            key = str(v).strip() if v not in (None, "") else "(пусто)"
            out[key] = out.get(key, 0) + 1
        return sorted(([k, n] for k, n in out.items()), key=lambda x: -x[1])[:12]

    uch = safe("SELECT lesnichestvo, god_sozdaniya FROM lesokultury_uchastok")
    uch_f = [r for r in uch if _lesn_match(target, r[0])]
    mer = safe("SELECT m.tip, m.data, u.lesnichestvo FROM lesokultury_meropriyatiya m "
               "JOIN lesokultury_uchastok u ON u.id = m.uchastok_id")
    items = safe("SELECT lesnichestvo, status_rabot, data_vypolneniya FROM delyanka_item")
    acts = safe("SELECT act_date FROM osvidetelstvovanie_acts")
    return {
        "lesnichestvo_filtr": lesnichestvo,
        "kultury_vsego": len(uch),
        "kultury_v_lesnichestve": len(uch_f),
        "kultury_lesnichestva": counter(r[0] for r in uch),
        "kultury_gody_sozdaniya": counter(_year_of(r[1]) or r[1] for r in uch_f),
        "zhurnal_po_godam": counter(f"{r[0]} — {_year_of(r[1]) or r[1] or '(без даты)'}"
                                    for r in mer if _lesn_match(target, r[2])),
        "delyanki_lesnichestva": counter(r[0] for r in items),
        "delyanki_statusy": counter(f"{r[1] or '(нет)'} — {_year_of(r[2]) or 'без даты'}"
                                    for r in items if _lesn_match(target, r[0])),
        "akty_po_godam": counter(_year_of(r[0]) or r[0] for r in acts),
    }


# Графы, без которых строка считается недописанной (подсвечиваются на сайте;
# когда они заполнены, замечания по строке больше не показываются).
NUZHNYE = {
    3: [2, 5, 6],
    4: [1, 2, 3, 4, 5, 6, 7, 8, 9],
    7: [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11],
    14: [2, 3],
    15: [2, 5, 6, 7],
}


def ensure_popravki(conn) -> None:
    conn.execute(
        "CREATE TABLE IF NOT EXISTS tek_izm_popravki ("
        " god INTEGER NOT NULL,"
        " prilozhenie INTEGER NOT NULL,"
        " klyuch TEXT NOT NULL,"
        " znacheniya_json TEXT NOT NULL,"
        " updated_at TEXT DEFAULT (datetime('now', 'localtime')),"
        " PRIMARY KEY (god, prilozhenie, klyuch))"
    )


def popravki(conn, god: int) -> Dict[Tuple[int, str], Dict[int, str]]:
    """Дописанное вручную к автоматическим строкам: {(прил, ключ): {графа: значение}}."""
    ensure_popravki(conn)
    out = {}
    for pril, key, values in conn.execute(
        "SELECT prilozhenie, klyuch, znacheniya_json FROM tek_izm_popravki WHERE god = ?", (god,)
    ).fetchall():
        try:
            parsed = json.loads(values)
        except ValueError:
            continue
        if isinstance(parsed, dict):
            out[(pril, key)] = {int(k): str(v) for k, v in parsed.items() if str(k).isdigit()}
    return out


def save_popravka(conn, god: int, prilozhenie: int, klyuch: str, values: List[str], lesnichestvo: str = "") -> None:
    """Сохраняет только графы, отличающиеся от автоматических значений
    (чтобы исправление данных в программе не перекрывалось старой правкой)."""
    ensure_popravki(conn)
    data = build(conn, god, lesnichestvo, s_popravkami=False)
    res = data.get(prilozhenie)
    if res is None or klyuch not in res["keys"]:
        raise TIError("Строка не найдена — обновите страницу")
    avto = res["rows"][res["keys"].index(klyuch)]
    diff = {str(i): str(v or "").strip() for i, v in enumerate(values[:len(avto)])
            if str(v or "").strip() != str(avto[i] or "").strip()}
    if diff:
        conn.execute(
            "INSERT INTO tek_izm_popravki (god, prilozhenie, klyuch, znacheniya_json) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(god, prilozhenie, klyuch) DO UPDATE SET znacheniya_json = excluded.znacheniya_json, "
            "updated_at = datetime('now', 'localtime')",
            (god, prilozhenie, klyuch, json.dumps(diff, ensure_ascii=False)),
        )
    else:
        conn.execute("DELETE FROM tek_izm_popravki WHERE god = ? AND prilozhenie = ? AND klyuch = ?",
                     (god, prilozhenie, klyuch))
    conn.commit()


# Столбец «новый выдел (подвыдел)» в приложениях, где он есть.
NOVYY_VYDEL_COL = {3: 4, 4: 3, 7: 3, 15: 4}


def _nomer(text) -> Optional[int]:
    """Целая часть номера выдела: «33» -> 33, «33.1» -> 33, «12а» -> 12."""
    m = re.match(r"\s*(\d+)", str(text or ""))
    return int(m.group(1)) if m else None


def _novye_vydely(result: dict, taxation: dict) -> None:
    """Каждому участку / лесосеке — номер, который можно дать новому выделу:
    последний выдел квартала по таксации + 1, пропуская номера, уже
    присвоенные в этом году другим участкам того же квартала. Сам номер
    ставит человек (кнопка «Новый выдел» / «По выделам»)."""
    maks: Dict[str, int] = {}
    for (kv, vd) in taxation:
        n = _nomer(vd)
        if n is not None:
            maks[kv] = max(maks.get(kv, 0), n)
    zanyato: Dict[str, Dict[str, set]] = {}
    for n, col in NOVYY_VYDEL_COL.items():
        res = result[n]
        for key, row in zip(res["keys"], res["rows"]):
            if not key or col >= len(row):
                continue
            kv = str(row[0] or "").strip()
            nomer = _nomer(row[col])
            if nomer is not None and str(row[col]).strip() != str(row[1] or "").strip():
                zanyato.setdefault(kv, {}).setdefault(str(key).split(":")[0], set()).add(nomer)
    for n in NOMERA:
        for prefix, info in result[n].get("chasti_info", {}).items():
            kv = str(info.get("kvartal") or "").strip()
            chuzhie = [x for p, nums in zanyato.get(kv, {}).items() if p != prefix for x in nums]
            svoi = sorted(zanyato.get(kv, {}).get(prefix, set()))
            # Без таксации квартала номер не угадать — пусть впишут сами.
            info["novyy"] = (str(svoi[0]) if svoi
                             else str(max([maks[kv]] + chuzhie) + 1) if kv in maks else "")
            info["posledniy"] = str(maks.get(kv, "")) if kv in maks else ""


def _key_matches(warn_key: str, row_key: str) -> bool:
    return row_key == warn_key or row_key.startswith(warn_key + ":")


def build(conn, god: int, lesnichestvo: str = "", s_popravkami: bool = True) -> dict:
    """Строки всех приложений за год:
    {n: {"rows": [[...]], "keys": [...], "pustye": [[графы]], "popravleno": [bool],
         "warnings": [...], "uchastki": k, "avto": m, "ruchnye": [{id, values}]}}
    и сводная прил. 2 в result["svodnaya"]."""
    taxation = _taxation_ploshad(conn, lesnichestvo)
    result = {n: {"rows": [], "keys": [], "warn_keys": [], "warnings": [], "uchastki": 0, "ruchnye": []}
              for n in NOMERA}
    _kultury(conn, god, lesnichestvo, result, taxation)
    _rubki_delyanki(conn, god, lesnichestvo, result, taxation)
    _rubki_uhoda(conn, god, lesnichestvo, result, taxation)
    _novye_vydely(result, taxation)
    popr = popravki(conn, god) if s_popravkami else {}
    for n in NOMERA:
        res = result[n]
        # Дописанное вручную поверх автоматических значений.
        res["popravleno"] = []
        for i, key in enumerate(res["keys"]):
            p = popr.get((n, key))
            res["popravleno"].append(bool(p))
            if not p:
                continue
            row = res["rows"][i]
            for col, value in p.items():
                if col < len(row):
                    row[col] = value
            if n in (3, 15) and 3 in p and res.get("vidy"):
                vid = vid_rubki(p[3])
                if vid and pril_rubki(vid) == n:
                    res["vidy"][i] = vid
            elif n == 7 and 5 in p and res.get("vidy"):
                res["vidy"][i] = _vid_kultur(p[5])
        # Недописанные графы и замечания только по недописанным строкам.
        nuzhnye = NUZHNYE.get(n, [])
        res["pustye"] = [[c for c in nuzhnye if c < len(r) and not str(r[c] or "").strip()] for r in res["rows"]]
        for key, text, cols in res.pop("warn_keys"):
            if key is not None:
                idx = [i for i, k in enumerate(res["keys"]) if _key_matches(key, k)]
                check = cols if cols is not None else nuzhnye
                if idx and all(str(res["rows"][i][c] or "").strip()
                               for i in idx for c in check if c < len(res["rows"][i])):
                    continue
            res["warnings"].append(text)
        # Авто-строки сортируем по кварталу/выделу, ручные идут следом в порядке ввода.
        vidy = res.get("vidy")
        res.setdefault("dop", [""] * len(res["rows"]))
        cols = [res["rows"], res["keys"], res["pustye"], res["popravleno"], res["dop"]]
        order = sorted(range(len(res["rows"])),
                       key=lambda i: (_sort_num(res["rows"][i][0]), _sort_num(res["rows"][i][1])))
        res["rows"], res["keys"], res["pustye"], res["popravleno"], res["dop"] = ([c[i] for i in order] for c in cols)
        if vidy is not None:
            res["vidy"] = [vidy[i] for i in order]
        res["avto"] = len(res["rows"])
    _ruchnye_v_result(conn, god, lesnichestvo, result)
    result["svodnaya"] = svodnaya({n: result[n] for n in NOMERA})
    return result


def _sort_num(text) -> tuple:
    parts = re.findall(r"\d+", str(text))
    return tuple(int(x) for x in parts) or (10 ** 9,)


# --------------------------------------------------------------------------- #
#   Word
# --------------------------------------------------------------------------- #
def _para_text(p) -> str:
    return "".join(r.text for r in p.runs)


def _set_para_text(p, text: str) -> None:
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)


def _set_cell(cell, text: str) -> None:
    paras = cell.paragraphs
    for extra in paras[1:]:
        extra._element.getparent().remove(extra._element)
    _set_para_text(paras[0], text)


def _is_data_row(row) -> bool:
    first = row.cells[0].text.strip()
    return first == "" or bool(re.fullmatch(r"\d+([.,]\d+)?", first))


def _fill_table(table, rows: List[List[str]], min_rows: int = 3) -> None:
    trs = table.rows
    data_idx = next((i for i, r in enumerate(trs) if _is_data_row(r)), len(trs))
    proto = copy.deepcopy((trs[data_idx] if data_idx < len(trs) else trs[-1])._tr)
    for r in list(trs)[data_idx:]:
        r._tr.getparent().remove(r._tr)
    for values in rows or [[""] * len(table.columns)] * min_rows:
        tr = copy.deepcopy(proto)
        table._tbl.append(tr)
        row = table.rows[-1]
        for i, cell in enumerate(row.cells):
            _set_cell(cell, values[i] if i < len(values) else "")


def _tables_by_prilozhenie(doc) -> Dict[int, object]:
    """Таблица каждого приложения — по ближайшему выше абзацу «Приложение N»."""
    from docx.table import Table
    result = {}
    current = None
    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = "".join(t.text or "" for t in el.iter() if t.tag.endswith("}t"))
            m = re.search(r"Приложение\s+(\d+)\s*$", text.strip()) or re.match(r"^\s*Приложение\s+(\d+)\b", text)
            if m and "к Инструкции" not in text:
                current = int(m.group(1))
        elif tag == "tbl" and current is not None and current not in result:
            result[current] = Table(el, doc._body)
    return result


def _update_headers(doc, god: int, data_zapolneniya: str,
                    ploshad_nachalo: str = "", ploshad_konec: str = "") -> None:
    for p in doc.paragraphs:
        text = _para_text(p)
        new = re.sub(r"(текущих изменений за\s*)(?:19|20)\d{2}", lambda m: f"{m.group(1)}{god}", text)
        new = re.sub(r"(Дата заполнения[_\s]*)\d{2}\.\d{2}\.\d{4}", lambda m: f"{m.group(1)}{data_zapolneniya}", new)
        # Прил. 1: «Общая площадь лесничества на 01.01.ГГГГ года ___ га (начало/конец отчетного года)».
        konec = "конец отчетного" in new
        new = re.sub(r"(на 01\.01\.)(?:19|20)\d{2}", lambda m: f"{m.group(1)}{god + 1 if konec else god}", new)
        ploshad = ploshad_konec if konec else ploshad_nachalo
        if ploshad and "Общая площадь" in new:
            new = re.sub(r"(года\s*)_+\s*(га)", lambda m: f"{m.group(1)}{ploshad} {m.group(2)}", new)
        if new != text:
            _set_para_text(p, new)


def _fill_svodnaya(table, svod: Dict[int, Tuple[float, int]]) -> None:
    rows = table.rows
    for row_no, (area, count) in svod.items():
        if row_no >= len(rows):
            continue
        cells = rows[row_no].cells
        _set_cell(cells[1], fmt(area) if count else "")
        _set_cell(cells[2], str(count) if count else "")


def make_docx(data: dict, god: int, template: Optional[bytes] = None,
              data_zapolneniya: Optional[str] = None,
              ploshad_nachalo: str = "", ploshad_konec: str = "") -> bytes:
    """Заполняет таблицы прил. 3–15 строками из build(), прил. 2 — сводной,
    в шапках меняет год и дату заполнения (в прил. 1 — и общую площадь)."""
    from docx import Document

    try:
        doc = Document(io.BytesIO(template) if template else str(TEMPLATE_PATH))
    except Exception as exc:  # noqa: BLE001
        raise TIError(f"Не удалось открыть шаблон Word: {exc}")
    tables = _tables_by_prilozhenie(doc)
    nomera = [n for n in data if isinstance(n, int)]
    missing = [n for n in nomera + [2] if n not in tables]
    if missing:
        raise TIError("В шаблоне не найдены таблицы приложений: " + ", ".join(str(n) for n in sorted(missing)))
    for n in nomera:
        _fill_table(tables[n], data[n]["rows"])
    if "svodnaya" in data:
        _fill_svodnaya(tables[2], data["svodnaya"])
    _update_headers(doc, god, data_zapolneniya or dt.date.today().strftime("%d.%m.%Y"),
                    str(ploshad_nachalo or "").strip(), str(ploshad_konec or "").strip())
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
